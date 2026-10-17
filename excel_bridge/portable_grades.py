"""Calculate and export complete grades using the original .xlsm, without Excel COM."""
import hashlib
import io
import json
import os
import shutil
import tempfile
from contextlib import contextmanager
from datetime import datetime

import openpyxl
from openpyxl.utils.cell import get_column_letter

from catalog import macro_signature
from formula_engine import FormulaEngine,ExcelError
from grade_repository import GradeRepository
from validation import validate_grades
from xlsm_package import patch_package,vba_hashes


def sha(data):return hashlib.sha256(data).hexdigest()


class PortableGrades:
    def __init__(self,catalog,database):
        self.catalog=catalog
        self.repository=GradeRepository(database)

    def key(self,class_id):return self.catalog.public()['year']+':'+class_id

    @contextmanager
    def lock(self,class_id):
        folder=self.catalog.state_dir/'grade-locks';folder.mkdir(parents=True,exist_ok=True)
        path=folder/(class_id+'.lock')
        with path.open('a+b') as f:
            if f.tell()==0:f.write(b'0');f.flush()
            f.seek(0)
            try:
                if os.name=='nt':
                    import msvcrt
                    msvcrt.locking(f.fileno(),msvcrt.LK_NBLCK,1)
                else:
                    import fcntl
                    fcntl.flock(f.fileno(),fcntl.LOCK_EX|fcntl.LOCK_NB)
            except OSError:raise ValueError('Lớp đang được cập nhật điểm ở phiên khác. Chờ hoàn tất rồi thử lại.')
            try:yield
            finally:
                f.seek(0)
                if os.name=='nt':msvcrt.locking(f.fileno(),msvcrt.LK_UNLCK,1)
                else:fcntl.flock(f.fileno(),fcntl.LOCK_UN)

    def info(self,class_id):
        self.catalog.refresh()
        info=self.catalog.classes.get(class_id)
        if not info:raise ValueError('Không tìm thấy lớp.')
        return info

    def model(self,package):
        book=openpyxl.load_workbook(io.BytesIO(package),read_only=True,data_only=False)
        try:return FormulaEngine(book)
        finally:book.close()

    def rows(self,engine,sheet,first,last):
        rows={}
        for row in range(first,last+1):
            value=engine.cell(sheet,f'B{row}')
            if not value:continue
            code=str(value).strip().upper()
            if code in rows:raise ValueError(f'Mã {code} trùng trong {sheet}.')
            rows[code]=row
        return rows

    def response(self,engine,info,semester,revision):
        name='HK '+str(semester) if semester in (1,2) else 'Tổng kết'
        first,last=(6,45) if semester in (1,2) else (7,46)
        results,issues=[],[]
        def value(ref):
            result=engine.cell(name,ref)
            if isinstance(result,ExcelError):
                issues.append({'sheet':name,'cell':ref,'error':result.code});return result.code
            return None if result=='' else result
        for code,row in self.rows(engine,name,first,last).items():
            person=self.catalog.students.get(code)
            if not person or person['classId']!=info['id']:continue
            student={'code':code,'name':person['name'],'saintName':person['saintName'],'fullName':person['fullName'],'rosterRow':person['rosterRow']}
            if semester in (1,2):
                student.update({'khaoBai':[value(f'{get_column_letter(c)}{row}') for c in range(5,25)],
                    **{k:value(f'{c}{row}') for k,c in {'kt15':'Y','kt1Tiet':'Z','thiKinh':'AA','thiHk':'AB',
                       'tbHk':'AC','xepLoai':'AD','xepHang':'AE','massRatio':'AF','lessonRatio':'AG'}.items()}})
            else:
                student.update({k:value(f'{c}{row}') for k,c in {'diemHk1':'E','diemHk2':'F','diemGiaoLy':'G',
                    'xepLoai':'H','massScore':'I','lessonScore':'J','combinedMassScore':'K',
                    'assessment':'L','totalScore':'M','yearRank':'N','supplementalMass':'O'}.items()})
            results.append(student)
        return {'classId':info['id'],'className':info['name'],'semester':semester,'students':results,
                'revision':revision,'calculatedBy':'server','calculationIssues':issues,
                'thresholds':{'catechism':engine.cell('Tổng kết','D3'),'attendance':engine.cell('Tổng kết','D4')}}

    def _replace(self,info,expected,package):
        source=info['path']
        owner_files=[source.with_name('~$'+source.name),source.with_name('~$'+source.name[2:])]
        if any(p.exists() for p in owner_files):
            raise PermissionError('File lớp có khóa của Excel. Đóng file trước khi cập nhật.')
        original=source.read_bytes()
        if sha(original)!=expected:raise ValueError('File lớp thay đổi bên ngoài. Chưa ghi đè; tải lại bảng điểm.')
        if vba_hashes(original)!=vba_hashes(package):raise ValueError('VBA trong file xuất không khớp mẫu.')
        backup=self.catalog.state_dir/'backups'/info['id'];backup.mkdir(parents=True,exist_ok=True)
        with tempfile.TemporaryDirectory(prefix='.grade-stage-',dir=source.parent) as folder:
            from pathlib import Path
            stage=Path(folder)/source.name
            stage.write_bytes(package)
            if macro_signature(stage)[0]!=self.catalog.trusted_signature:
                raise ValueError('Mã nguồn VBA thay đổi. Chưa cập nhật sổ lớp.')
            if sha(source.read_bytes())!=expected:raise ValueError('Excel đã thay đổi trong lúc tính điểm. Chưa thay file gốc.')
            destination=backup/(datetime.now().strftime('%Y%m%d-%H%M%S-%f')+'.xlsm')
            shutil.copy2(source,destination)
            if sha(destination.read_bytes())!=expected or sha(source.read_bytes())!=expected:
                raise ValueError('Sổ thay đổi trong lúc sao lưu. Chưa thay thế.')
            os.replace(stage,source)
        self.catalog.refresh()

    def recover(self,info):
        key=self.key(info['id']);job=self.repository.job(key)
        if job and job['status']=='conflict':
            raise ValueError('Có điểm chờ xuất bị xung đột với sửa Excel bên ngoài. Chưa ghi đè file; cần đối chiếu lịch sử.')
        if not job or job['status']!='pending':return None
        current=sha(info['path'].read_bytes())
        if current==job['result_sha']:
            self.repository.complete(key);return None
        if current!=job['expected_sha']:
            self.repository.forget_conflict(key)
            raise ValueError('File được sửa bên ngoài khi còn điểm chờ xuất. Kiểm tra lịch sử trước khi nhập tiếp.')
        package=bytes(job['package'])
        if sha(package)!=job['result_sha']:raise ValueError('Bản Excel chờ xuất không khớp mã kiểm tra.')
        self._replace(info,current,package)
        self.repository.complete(key)
        return None

    def read(self,class_id,semester=1):
        if type(semester) is not int or semester not in (1,2,3):raise ValueError('Học kỳ không hợp lệ.')
        info=self.info(class_id)
        job=self.repository.job(self.key(class_id))
        if job and job['status']=='conflict':
            result=self.response(self.model(info['path'].read_bytes()),info,semester,sha(info['path'].read_bytes()))
            return {**result,'excelReady':False,'exportConflict':True}
        if job and job['status']=='pending':
            # Compute pending packages too, including new summary columns after an upgrade.
            engine=self.model(bytes(job['package']));engine.grade_outputs()
            result=self.response(engine,info,semester,job['result_sha'])
            return {**result,'excelReady':False}
        package=info['path'].read_bytes();engine=self.model(package)
        engine.grade_outputs()
        return {**self.response(engine,info,semester,sha(package)),'excelReady':True}

    def save(self,class_id,semester,grades,expected_revision=None):
        if type(semester) is not int or semester not in (1,2):raise ValueError('Học kỳ phải là 1 hoặc 2.')
        info=self.info(class_id)
        grades=validate_grades(grades,class_id,self.catalog.students)
        with self.lock(class_id):
            self.recover(info)
            package=info['path'].read_bytes();before=sha(package)
            if expected_revision is not None and expected_revision!=before:
                raise ValueError('Bảng điểm đã được sửa ở phiên khác. Tải lại trước khi lưu để tránh ghi đè.')
            engine=self.model(package);sheet='HK '+str(semester);rows=self.rows(engine,sheet,6,45)
            inputs={}
            for item in grades:
                row=rows.get(item['code'])
                if row is None:raise ValueError('Mã không còn trong sheet đích: '+item['code'])
                for index,score in enumerate(item.get('khaoBai',[])):inputs[(sheet,f'{get_column_letter(5+index)}{row}')]=score
                for field,col in {'kt15':'Y','kt1Tiet':'Z','thiKinh':'AA','thiHk':'AB'}.items():
                    if field in item:inputs[(sheet,f'{col}{row}')]=item[field]
            for (s,ref),score in inputs.items():
                old=engine.cells.get((s,ref))
                if isinstance(old,str) and old.startswith('='):raise ValueError(f'Ô {s}!{ref} có công thức. Chưa ghi đè.')
                engine.set(s,ref,score)
            caches=engine.grade_outputs();result=patch_package(package,inputs,caches);after=sha(result)
            books={str(s):self.response(engine,info,s,after) for s in (1,2,3)}
            self.repository.enqueue(self.key(class_id),before,after,result,{'changes':{'semester':semester,'grades':grades},'book':books})
            ready=True;conflict=False
            try:self.recover(info)
            except OSError:ready=False
            except ValueError:
                ready=False
                current=sha(info['path'].read_bytes())
                if current not in (before,after):
                    self.repository.forget_conflict(self.key(class_id));conflict=True
            message='Đã lưu điểm và kết quả vào Excel.' if ready else 'Điểm đã được lưu; Excel chưa cập nhật. Đóng file nếu đang mở và kiểm tra quyền ghi/dung lượng, rồi tải lại.'
            if conflict:message='Điểm đã được lưu; file Excel vừa được sửa bên ngoài. Chưa ghi đè, cần đối chiếu bản chờ với file hiện tại.'
            return {'saved':True,'excelReady':ready,'vbaExecuted':False,'calculatedBy':'server',
                    'grades':{**books[str(semester)],'excelReady':ready,'exportConflict':conflict,
                              'revision':sha(info['path'].read_bytes()) if conflict else after},'revision':after,'message':message}

    def export(self,class_id):
        info=self.info(class_id)
        with self.lock(class_id):
            self.recover(info)
            package=info['path'].read_bytes();engine=self.model(package)
            result=patch_package(package,{},engine.grade_outputs())
            if result!=package:self._replace(info,sha(package),result)
            return result,info['filename'],'application/vnd.ms-excel.sheet.macroEnabled.12'

    def status(self):return 'mysql' if self.repository.mysql else 'sqlite'

    def pending_export(self,class_id):
        info=self.info(class_id);job=self.repository.job(self.key(class_id))
        if not job or job['status'] not in ('pending','conflict') or not job['package']:
            raise ValueError('Không có bản điểm chờ để đối chiếu.')
        package=bytes(job['package'])
        if sha(package)!=job['result_sha']:raise ValueError('Bản chờ không khớp mã kiểm tra.')
        return package,'DiemChoDoiChieu_'+info['filename'],'application/vnd.ms-excel.sheet.macroEnabled.12'

    def resolve_conflict(self,class_id,values):
        if not values.get('_manager'):raise ValueError('Người quản lý trên máy mới được đối chiếu điểm chờ xuất.')
        actor=str(values.get('actor','')).strip();reason=str(values.get('reason','')).strip()
        if not 2<=len(actor)<=100 or not 3<=len(reason)<=500:raise ValueError('Nhập tên người xử lý và lý do đối chiếu.')
        info=self.info(class_id)
        with self.lock(class_id):
            job=self.repository.job(self.key(class_id))
            if not job or job['status']!='conflict':raise ValueError('Không còn xung đột điểm cần đối chiếu.')
            current=info['path'].read_bytes()
            if values.get('expectedRevision')!=sha(current):raise ValueError('File đã thay đổi. Làm mới bảng điểm trước khi đối chiếu.')
            if values.get('decision')!='keep-file':raise ValueError('Chỉ xác nhận giữ file hiện tại; nhập lại điểm cần giữ sau khi đối chiếu.')
            # Preserve the pending workbook before explicitly choosing the externally edited file.
            archive=self.catalog.state_dir/'grade-conflicts'/class_id
            archive.mkdir(parents=True,exist_ok=True)
            stamp=datetime.now().strftime('%Y%m%d-%H%M%S-%f')
            (archive/(stamp+'.xlsm')).write_bytes(bytes(job['package']))
            (archive/(stamp+'.json')).write_text(json.dumps({'actor':actor,'reason':reason,'changes':json.loads(job['inputs'])['changes']},ensure_ascii=False),encoding='utf-8')
            engine=self.model(current);engine.grade_outputs()
            book={str(s):self.response(engine,info,s,sha(current)) for s in (1,2,3)}
            self.repository.accept_file(self.key(class_id),sha(current),book,actor,reason)
            return {'saved':True,'grades':book[str(values.get('semester',1))],
                    'message':'Đã giữ file hiện tại. Bản điểm chờ xuất được lưu riêng để đối chiếu và nhập lại nếu cần.'}
