"""Read the supplied class workbooks without rewriting their OOXML/VBA."""
import hashlib
import json
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

import openpyxl
from oletools.olevba import VBA_Parser

ACTIVITIES = {'mass': range(5, 19), 'early': range(19, 24), 'lesson': range(24, 29)}
SAINTS = set('maria giuse teresa anna phero gioan đaminh phaolo martino vincente daminh phanxico gioakim anton matta cecilia micae catarina f.xavie augustino rosa simon phalo laureso tomaso macco miacae đaminh gioan.b toma anre gioankim gerado têresa emmanual giacobe don'.split())
MULTI_SAINTS = {('gioan', 'b'), ('gioan', 'baotixita'), ('gioan', 'bosco'), ('maria', 'madalena'), ('maria', 'teresa'), ('don', 'bosco')}


def clean(value):
    return ' '.join(unicodedata.normalize('NFC', str(value or '')).split())


def split_name(family, given):
    words = clean(family).split()
    prefix = 0
    if words and words[0].casefold() in SAINTS:
        prefix = 2 if tuple(w.casefold() for w in words[:2]) in MULTI_SAINTS else 1
    return ' '.join(words[:prefix]), clean(' '.join(words[prefix:]) + ' ' + clean(given))


def iso_date(value):
    return value.date().isoformat() if isinstance(value, datetime) else value.isoformat() if isinstance(value, date) else None


def macro_signature(path):
    parser = VBA_Parser(str(path))
    try:
        modules = sorted((name, code) for _, _, name, code in parser.extract_macros())
        digest = hashlib.sha256('\n'.join(name + '\n' + code for name, code in modules).encode()).hexdigest()
        module, source = next((name, code) for name, code in modules if name == 'Module1.bas')
        proc = re.search(r'^Sub\s+([^\s(]+)', source, re.M).group(1)
        return digest, module.removesuffix('.bas') + '.' + proc
    finally:
        parser.close()


class Catalog:
    def __init__(self, root, state_dir, trusted_signature):
        self.root = Path(root).resolve()
        self.state_dir = Path(state_dir)
        self.trusted_signature = trusted_signature
        self.classes = {}
        self.students = {}
        self.marks = {}
        self.stamps = {}
        self.file_cache = {}
        self.refresh()

    def paths(self):
        return [p for p in sorted(self.root.glob('*.xlsm')) + sorted((self.root / 'outputs/cac-lop-2026-2027').glob('*.xlsm'))
                if not p.name.startswith('~$')]

    def refresh(self):
        paths = self.paths()
        if not paths:
            raise ValueError('Không tìm thấy file .xlsm của các lớp trong thư mục dữ liệu.')
        stamps = {str(p): (p.stat().st_mtime_ns, p.stat().st_size) for p in paths}
        if stamps == self.stamps:
            return
        classes, students, marks = {}, {}, {}
        for path in paths:
            cached = self.file_cache.get(str(path))
            if cached and self.stamps.get(str(path)) == stamps[str(path)]:
                info, people, previous_marks = cached
                if info['id'] in classes or any(s['id'] in students for s in people):
                    raise ValueError('Danh sách lớp hoặc mã học sinh bị trùng.')
                classes[info['id']] = info
                students.update({s['id']: s for s in people})
                marks.update(previous_marks)
                continue
            signature, macro = macro_signature(path)
            if signature != self.trusted_signature:
                raise ValueError(f'Macro trong {path.name} đã thay đổi. Cần kiểm tra trước khi chạy.')
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                roster = book['Danh Sách']
                class_name = clean(roster['B2'].value) + ' ' + clean(roster['E2'].value)
                class_id = hashlib.sha256(class_name.encode()).hexdigest()[:16]
                if 'XH_ClassId' in book.custom_doc_props.names:
                    class_id = str(book.custom_doc_props['XH_ClassId'].value)
                    if not re.fullmatch(r'[0-9a-f]{16}', class_id):
                        raise ValueError('Mã định danh lớp không hợp lệ.')
                if class_id in classes:
                    raise ValueError(f'Hai file cùng lớp {class_name}.')
                teachers = {
                    'glv1': clean(roster['B3'].value) or '',
                    'glv2': clean(roster['E3'].value) or '',
                    'souer': clean(roster['B4'].value) or '',
                    'glv3': clean(roster['E4'].value) or ''
                }
                profiles = {}
                if 'Lý Lịch' in book.sheetnames:
                    for row in book['Lý Lịch'].iter_rows(min_row=5, max_row=46, max_col=12, values_only=True):
                        if not row[1]:
                            continue
                        p_code = clean(row[1]).upper()
                        profiles[p_code] = {
                            'gender': clean(row[4]) or '',
                            'birthDate': str(row[5])[:10] if row[5] else '',
                            'baptismDate': str(row[6])[:10] if row[6] else '',
                            'parents': clean(row[7]) or '',
                            'address': clean(row[8]) or '',
                            'subParish': clean(row[9]) or '',
                            'phone': clean(row[10]) or ''
                        }
                people = []
                for rowno, row in enumerate(roster.iter_rows(min_row=7, max_row=46, max_col=10, values_only=True), 7):
                    if not row[1]:
                        continue
                    code = clean(row[1]).upper()
                    if not re.fullmatch(r'XH\d+', code) or code in students:
                        raise ValueError(f'Mã trùng hoặc sai định dạng: {code}.')
                    saint, full = split_name(row[2], row[3])
                    property_name = 'XH_Saint_' + code
                    if property_name in book.custom_doc_props.names:
                        explicit_saint = clean(book.custom_doc_props[property_name].value)
                        if explicit_saint=='__XH_EMPTY__':explicit_saint=''
                        family = clean(row[2])
                        if not explicit_saint or family == explicit_saint or family.startswith(explicit_saint + ' '):
                            saint = explicit_saint
                            full = clean(family[len(saint):] + ' ' + clean(row[3]))
                    prof = profiles.get(code, {})
                    person = {'id': code, 'saintName': saint, 'fullName': full, 'name': clean(saint + ' ' + full),
                              'className': class_name, 'classId': class_id, 'enrolledFrom': iso_date(row[8]),
                              'rosterRow': rowno,
                              'gender': prof.get('gender') or clean(row[9]) or '',
                              'birthDate': prof.get('birthDate') or '',
                              'baptismDate': prof.get('baptismDate') or '',
                              'parents': prof.get('parents') or '',
                              'address': prof.get('address') or '',
                              'subParish': prof.get('subParish') or '',
                              'phone': prof.get('phone') or ''}
                    students[code] = person
                    people.append(person)
                calendar, targets = {a: [] for a in ACTIVITIES}, {}
                for month in ['09', '10', '11', '12', '01', '02', '03', '04', '05', '06', '07', '08']:
                    sheet = book['T' + month]
                    rows = list(sheet.iter_rows(min_row=7, max_row=47, max_col=28, values_only=True))
                    for activity, columns in ACTIVITIES.items():
                        for col in columns:
                            day = iso_date(rows[0][col-1])
                            if not day:
                                continue
                            calendar[activity].append(day)
                            targets[(day, activity)] = ('T' + month, col)
                            for rowno, row in enumerate(rows[1:], 8):
                                code = clean(row[1]).upper()
                                status = clean(row[col-1]).upper()
                                if code in students and students[code]['classId'] == class_id and status in ('X', 'O', 'P'):
                                    marks[f'{day}|{activity}|{code}'] = {'status': status, 'timestamp': None, 'source': 'Excel'}
                semester_months={'hk1':[],'hk2':[]}
                for option_row in book['OPTION'].iter_rows(min_row=4,max_row=15,max_col=2,values_only=True):
                    option_date=iso_date(option_row[0])
                    semester=clean(option_row[1]).upper()
                    if option_date and semester in ('HKI','HKII'):
                        semester_months['hk1' if semester=='HKI' else 'hk2'].append(option_date[:7])
                classes[class_id] = {'id': class_id, 'name': class_name, 'filename': path.name, 'path': path,
                                     'count': len(people), 'calendar': calendar, 'targets': targets, 'macro': macro,
                                     'teachers': teachers, 'semesterMonths':semester_months}
                codes = {s['id'] for s in people}
                self.file_cache[str(path)] = (classes[class_id], people, {k: v for k, v in marks.items() if k.rsplit('|', 1)[-1] in codes})
            finally:
                book.close()
        self.classes, self.students, self.marks, self.stamps = classes, students, marks, stamps

    def eligible(self, session):
        return [s for s in self.students.values() if (session['scope'] == 'all' or s['className'] == session['scope'])
                and session['date'] in self.classes[s['classId']]['calendar'][session['activity']]
                and (not s['enrolledFrom'] or s['enrolledFrom'] <= session['date'])]

    def public(self):
        dates=sorted(d for c in self.classes.values() for d in c['calendar']['mass'])
        year=dates[0][:4] if dates else str(date.today().year)
        plans = {}
        source = Path(__file__).resolve().parents[1]/'src/data/staff-plan-2026-2027.json'
        if source.is_file():
            plan = json.loads(source.read_text(encoding='utf-8'))
            if plan.get('year') == f'{year}-{int(year)+1}':
                plans = {r['name']: {'location':r.get('location',''),'futureTeachers':r.get('futureTeachers',[])} for r in plan['classes']}
        return {'year':f'{year}-{int(year)+1}', 'students': list(self.students.values()), 'classes': [
            {**{k: c[k] for k in ('id', 'name', 'filename', 'count', 'calendar', 'teachers','semesterMonths')},
             'staffPlan':plans.get(c['name'])} for c in self.classes.values()],
            'marks': self.marks}

    def read_grades(self, class_id, semester=1):
        info = self.classes.get(class_id)
        if not info:
            raise ValueError('Không tìm thấy lớp.')
        sheet_name = 'HK 1' if semester == 1 else 'HK 2' if semester == 2 else 'Tổng kết'
        book = openpyxl.load_workbook(info['path'], read_only=True, data_only=True)
        try:
            if sheet_name not in book.sheetnames:
                raise ValueError(f'Không tìm thấy bảng điểm {sheet_name}.')
            sheet = book[sheet_name]
            class_students = [s for s in self.students.values() if s['classId'] == class_id]
            student_by_code = {s['id']: s for s in class_students}
            results = []
            if semester in (1, 2):
                for row in sheet.iter_rows(min_row=6, max_row=45, max_col=35, values_only=True):
                    if not row[1]:
                        continue
                    code = clean(row[1]).upper()
                    st = student_by_code.get(code)
                    if not st:
                        continue
                    khao_bai = [row[i] for i in range(4, 24)]
                    results.append({
                        'code': code,
                        'name': st['name'],
                        'saintName': st['saintName'],
                        'fullName': st['fullName'],
                        'rosterRow': st['rosterRow'],
                        'khaoBai': khao_bai,
                        'kt15': row[24],
                        'kt1Tiet': row[25],
                        'thiKinh': row[26],
                        'thiHk': row[27],
                        'tbHk': row[28],
                        'xepLoai': clean(row[29]),
                        'xepHang': row[30],
                        'massRatio': clean(row[31]),
                        'lessonRatio': clean(row[32])
                    })
            else:
                for row in sheet.iter_rows(min_row=7, max_row=46, max_col=15, values_only=True):
                    if not row[1]:
                        continue
                    code = clean(row[1]).upper()
                    st = student_by_code.get(code)
                    if not st:
                        continue
                    results.append({
                        'code': code,
                        'name': st['name'],
                        'saintName': st['saintName'],
                        'fullName': st['fullName'],
                        'rosterRow': st['rosterRow'],
                        'diemHk1': row[4],
                        'diemHk2': row[5],
                        'diemGiaoLy': row[6],
                        'xepLoai': clean(row[7]),
                        'massScore': row[8],
                        'lessonScore': row[9],
                        'combinedMassScore': row[10],
                        'assessment': row[11],
                        'totalScore': row[12],
                        'yearRank': row[13],
                        'supplementalMass': row[14]
                    })
            return {'classId': class_id, 'className': info['name'], 'semester': semester, 'students': results,
                    'thresholds':{'catechism':book['Tổng kết']['D3'].value,'attendance':book['Tổng kết']['D4'].value}}
        finally:
            book.close()
