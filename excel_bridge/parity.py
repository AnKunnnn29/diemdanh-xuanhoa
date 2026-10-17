"""Excel parity workflows with source previews and recoverable file + SQLite commits."""
import hashlib
import json
import os
import re
import secrets
import shutil
import sqlite3
from calendar import monthrange
from datetime import date, datetime
from pathlib import Path

import openpyxl
from catalog import clean, macro_signature
from development import digest
from exceptions import day
from validation import PROFILE_FIELDS, profile_date


def token(value):
    return hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True,default=str).encode()).hexdigest()


class ExcelParity:
    def init_parity(self):
        with self.db() as db:
            db.executescript('''
              CREATE TABLE IF NOT EXISTS issued_codes(code TEXT PRIMARY KEY, stamp TEXT);
              CREATE TABLE IF NOT EXISTS supplemental_mass(id TEXT PRIMARY KEY, code TEXT, actual_day TEXT,
                actor TEXT, reason TEXT, status TEXT, stamp TEXT, updated_at TEXT);
              CREATE UNIQUE INDEX IF NOT EXISTS one_supplement_per_day ON supplemental_mass(code,actual_day) WHERE status='active';
            ''')
            db.executemany('INSERT OR IGNORE INTO issued_codes VALUES (?,?)',[(s,self.stamp()) for s in self.catalog.students])
            for table in ('audit','scans','requests','student_status'):
                col='mark_key' if table=='audit' else 'code'
                for row in db.execute(f'SELECT DISTINCT {col} FROM {table}').fetchall():
                    code=(row[0] or '').rsplit('|',1)[-1]
                    if re.fullmatch(r'XH\d+',code):db.execute('INSERT OR IGNORE INTO issued_codes VALUES (?,?)',(code,self.stamp()))
            manifest=self.catalog.root/'manifest.json'
            if manifest.exists():
                previous=json.loads(manifest.read_text(encoding='utf-8'))
                codes=previous.get('issuedCodes',[])+[p.get('code','') for p in previous.get('placements',[])]
                for code in codes:
                    if isinstance(code,str) and re.fullmatch(r'XH\d+',code):db.execute('INSERT OR IGNORE INTO issued_codes VALUES (?,?)',(code,self.stamp()))

    def parity_engine(self):
        if self.engine is None:
            from server import ExcelEngine
            self.engine=ExcelEngine(self.catalog)
        return self.engine

    def parity_manager(self,values):
        if not values.get('_manager'):raise ValueError('Thực hiện thao tác quản lý sổ trên máy Excel.')
        return self.actor(values)

    def parity_guard(self,class_ids):
        self.catalog.refresh()
        for identifier in class_ids:
            if identifier not in self.catalog.classes:raise ValueError('Không tìm thấy lớp.')
            if self.grade_mode=='portable':
                service=self.grade_service()
                with service.lock(identifier):service.recover(self.catalog.classes[identifier])

    def parity_read(self,class_id):
        self.catalog.refresh()
        info=self.catalog.classes.get(class_id)
        if not info:raise ValueError('Không tìm thấy lớp.')
        book=openpyxl.load_workbook(info['path'],read_only=True,data_only=False)
        try:
            option=book['OPTION']
            schedules=[]
            for row in range(3,18):
                label=clean(option.cell(row,7).value)
                if label:schedules.append({'row':row,'label':label,'massDays':[d for d in range(7) if option.cell(row,8+d).value==1],
                                            'lessonDays':[d for d in range(7) if option.cell(row,17+d).value==1]})
            notes=[]
            sheet=book['Ghi Chú']
            last=1
            for row in sheet.iter_rows(max_row=min(sheet.max_row,1000),max_col=7):
                for c in row:
                    if c.value is None:continue
                    last=max(last,c.row)
                    if c.row>1:notes.append({'cell':c.coordinate,'content':str(c.value),'editable':isinstance(c.value,str) and not c.value.startswith('=')})
            with self.db() as db:
                supplements=[dict(r) for r in db.execute('SELECT * FROM supplemental_mass WHERE code IN (SELECT code FROM issued_codes) ORDER BY actual_day DESC')
                             if r['code'] in self.catalog.students and self.catalog.students[r['code']]['classId']==class_id]
                used={r[0] for r in db.execute('SELECT code FROM issued_codes')}
            codes=[int(s[2:]) for s in used if re.fullmatch(r'XH\d{5}',s)]
            suggestion='XH'+str(max(codes,default=0)+1).zfill(5) if codes and all(re.fullmatch(r'XH\d{5}',s) for s in used) else ''
            year=sorted(d for c in self.catalog.classes.values() for d in c['calendar']['mass'])
            return {'classId':class_id,'className':info['name'],'revision':digest(info['path']),
                    'configuration':{'semesters':[option.cell(r,2).value for r in range(4,16)],
                        'allowMass':option['E3'].value,'allowLesson':option['E4'].value,
                        'catechismThreshold':book['Tổng kết']['D3'].value,'attendanceThreshold':book['Tổng kết']['D4'].value,
                        'schedules':schedules},'notes':notes,'nextNoteCell':f'A{max(3,last+1)}',
                    'supplements':supplements,'suggestedCode':suggestion,
                    'yearStart':year[0][:4]+'-09-01' if year else '',
                    'history':[dict(r) for r in self._parity_history(class_id)]}
        finally:book.close()

    def _parity_history(self,class_id):
        with self.db() as db:
            return db.execute("SELECT * FROM management_events WHERE entity IN ('supplemental','enrollment','configuration','note','scratch','student-add') ORDER BY id DESC LIMIT 100").fetchall()

    def parity_dates(self,value):
        d=day(value)
        dates=sorted(x for c in self.catalog.classes.values() for x in c['calendar']['mass'])
        start=date(int(dates[0][:4]),9,1).isoformat();end=date(int(dates[0][:4])+1,8,31).isoformat()
        if not start<=d<=end:raise ValueError('Ngày phải trong năm học đang sử dụng.')
        return d

    def parity_plan(self,values):
        actor,reason=self.parity_manager(values)
        operation=values.get('operation');class_id=values.get('classId')
        self.parity_guard([class_id])
        info=self.catalog.classes[class_id]
        if values.get('expectedRevision') and values['expectedRevision']!=digest(info['path']):
            raise ValueError('Sổ lớp đã thay đổi. Làm mới thông tin trước khi nhập lại.')
        payload={};description=[];issues=[];identifiers=[class_id]
        if operation=='student-add':
            code=clean(values.get('code')).upper()
            if not re.fullmatch(r'XH\d+',code):raise ValueError('Mã phải có dạng XH và chữ số, dùng mã đã được xác nhận cấp.')
            with self.db() as db:
                if code in self.catalog.students or db.execute('SELECT 1 FROM issued_codes WHERE code=?',(code,)).fetchone():
                    raise ValueError('Mã đã được cấp hoặc có lịch sử. Không tái cấp mã cũ.')
            if info['count']>=40:raise ValueError('Lớp đã đủ 40 em.')
            payload={k:clean(values.get(k)) for k in PROFILE_FIELDS|{'saintName','fullName'}}
            if not 2<=len(payload['fullName'])<=150 or len(payload['saintName'])>100:raise ValueError('Nhập họ tên từ 2 đến 150 ký tự, tên thánh tối đa 100 ký tự.')
            if any(len(v)>500 for v in payload.values()):raise ValueError('Thông tin hồ sơ quá dài.')
            if payload['gender'] not in ('','Nam','Nữ'):raise ValueError('Giới tính phải là Nam hoặc Nữ.')
            for k in ('birthDate','baptismDate'):profile_date(payload[k])
            enrollment=self.parity_dates(values.get('enrolledFrom'))
            if enrollment[-2:]!='01':raise ValueError('Mẫu Excel quản lý nhập học theo tháng. Chọn ngày đầu tháng nhập học.')
            payload.update(code=code,enrolledFrom=enrollment)
            description=[f'Thêm {code} · {payload["fullName"]} vào {info["name"]}.',f'Nhập học từ tháng {enrollment[:7]}; giữ mã đã xác nhận cấp.',
                         'Cập nhật Danh Sách, Lý Lịch và kiểm tra liên kết 12 tháng, HK1, HK2, Tổng kết.']
        elif operation=='enrollment':
            code=clean(values.get('code')).upper();student=self.catalog.students.get(code)
            if not student or student['classId']!=class_id:raise ValueError('Thiếu nhi không thuộc lớp.')
            enrollment=self.parity_dates(values.get('enrolledFrom'))
            if enrollment[-2:]!='01':raise ValueError('Chọn ngày đầu tháng nhập học.')
            before=student['enrolledFrom']
            if enrollment==before:raise ValueError('Tháng nhập học chưa thay đổi.')
            marked=[k for k in self.catalog.marks if k.endswith('|'+code) and k.split('|')[0]<enrollment]
            if marked:issues.append('Có chuyên cần trước tháng nhập học mới; giữ nguyên lịch sử, chưa đổi tháng.')
            with self.db() as db:
                if db.execute("SELECT 1 FROM requests WHERE code=? AND status IN ('pending','approved','completed') AND (day<? OR COALESCE(exception_day,target_day)<?)",(code,enrollment,enrollment)).fetchone():
                    issues.append('Có đơn trước tháng nhập học mới. Xử lý đơn trước khi đổi.')
                if db.execute("SELECT 1 FROM supplemental_mass WHERE code=? AND status='active' AND actual_day<?",(code,enrollment)).fetchone():issues.append('Có lễ bổ sung trước tháng nhập học mới.')
            # Changing HK eligibility with entered points needs explicit separate reconciliation.
            old_sem='hk1' if not before or before[:7] in info['semesterMonths']['hk1'] else 'hk2'
            new_sem='hk1' if enrollment[:7] in info['semesterMonths']['hk1'] else 'hk2'
            if old_sem!=new_sem:
                for sem in (1,2):
                    r=next((r for r in self.read_grades(class_id,sem)['students'] if r['code']==code),None)
                    if r and any(v not in (None,'') for v in r['khaoBai']+[r['kt15'],r['kt1Tiet'],r['thiKinh'],r['thiHk']]):issues.append('Đổi học kỳ nhập học khi đã có điểm cần đối chiếu riêng; chưa áp dụng.');break
            payload={'code':code,'enrolledFrom':enrollment}
            description=[f'{code}: {before or "Cả năm"} → {enrollment[:7]}.','Chuyên cần và công thức học kỳ sẽ tính từ tháng nhập học mới.']
        elif operation=='supplemental':
            code=clean(values.get('code')).upper();student=self.catalog.students.get(code)
            if not student or student['classId']!=class_id:raise ValueError('Thiếu nhi không thuộc lớp.')
            existing_id=values.get('id');action=values.get('action','add')
            with self.db() as db:
                existing=db.execute('SELECT * FROM supplemental_mass WHERE id=?',(existing_id,)).fetchone() if existing_id else None
            if action not in ('add','cancel'):raise ValueError('Chỉ thêm hoặc hủy lượt; hủy rồi lập lại để sửa ngày.')
            if action=='cancel':
                if not existing or existing['code']!=code or existing['status']!='active':raise ValueError('Lượt lễ bổ sung không còn hiệu lực.')
                actual=existing['actual_day']
            else:
                actual=self.parity_dates(values.get('actualDate'))
                if actual>datetime.now(self.local_tz).date().isoformat():raise ValueError('Không ghi nhận lễ chưa diễn ra.')
                if not self.student_active(code,actual) or student['enrolledFrom'] and actual<student['enrolledFrom']:raise ValueError('Em chưa nhập học hoặc đang nghỉ vào ngày này.')
                if actual in info['calendar']['mass']:raise ValueError('Ngày này có lễ theo lịch lớp. Dùng điểm danh/đổi lễ, không cộng bổ sung.')
                with self.db() as db:
                    if db.execute("SELECT 1 FROM supplemental_mass WHERE code=? AND actual_day=? AND status='active'",(code,actual)).fetchone():raise ValueError('Đã ghi lễ bổ sung ngày này.')
                    if db.execute("SELECT 1 FROM scans WHERE code=? AND activity='mass' AND actual_day=?",(code,actual)).fetchone():raise ValueError('Lượt đi lễ đã dùng cho điểm danh/đổi lễ.')
                    if db.execute("SELECT 1 FROM requests WHERE code=? AND activity='mass' AND status IN ('pending','approved','completed') AND (day=? OR COALESCE(exception_day,target_day)=? OR actual_day=?)",(code,actual,actual,actual)).fetchone():raise ValueError('Ngày đi lễ đang có đơn hoặc đã dùng cho đổi lễ.')
            total=self.read_grades(class_id,3)
            current=next(r for r in total['students'] if r['code']==code)['supplementalMass'] or 0
            if isinstance(current,bool) or not isinstance(current,(int,float)) or current<0 or int(current)!=current:raise ValueError('Số lễ bù cũ không hợp lệ. Cần đối chiếu sổ.')
            after=current+(1 if action=='add' else -1)
            if after<0:raise ValueError('Số lễ bù trong sổ đã thay đổi; không thể trừ thêm.')
            payload={'code':code,'actualDate':actual,'action':action,'id':existing_id,'before':current,'after':after}
            tracked=sum(r['status']=='active' for r in self.parity_read(class_id)['supplements'] if r['code']==code)
            if current<tracked:raise ValueError('Số lễ bổ sung trong Excel nhỏ hơn các lượt đang có hiệu lực. Cần đối chiếu trước khi thêm/hủy.')
            description=[f'{code} · lễ thực tế {actual}: {"thêm" if action=="add" else "hủy"} 1 buổi.',f'Cột Đi Lễ Bù: {current} → {after}.',
                         f'{max(0,current-tracked)} buổi đã có trong sổ chưa có lịch sử từng ngày; giữ nguyên phần này.',
                         'Tính lại điểm chuyên cần, đánh giá, tổng điểm và hạng; không ghi thêm X vào ngày khác.']
        elif operation=='configuration':
            config=values.get('configuration')
            if not isinstance(config,dict):raise ValueError('Cấu hình không hợp lệ.')
            current=self.parity_read(class_id)['configuration']
            semesters=config.get('semesters');schedules=config.get('schedules')
            if not isinstance(semesters,list) or len(semesters)!=12 or any(v not in ('HKI','HKII') for v in semesters):raise ValueError('Chọn học kỳ cho đủ 12 tháng.')
            if 'HKI' not in semesters or 'HKII' not in semesters:raise ValueError('Năm học phải có cả hai học kỳ.')
            cells={f'B{i+4}':v for i,v in enumerate(semesters)}
            for key,ref in (('allowMass','E3'),('allowLesson','E4')):
                v=config.get(key)
                if type(v) is not int or not 0<=v<=31:raise ValueError('Số buổi được nghỉ phải là số nguyên 0–31.')
                cells[ref]=v
            thresholds={}
            for key,ref in (('catechismThreshold','D3'),('attendanceThreshold','D4')):
                v=config.get(key)
                if isinstance(v,bool) or not isinstance(v,(int,float)) or not 0<=v<=10:raise ValueError('Ngưỡng đạt phải từ 0 đến 10.')
                thresholds[ref]=v
            expected={s['row'] for s in current['schedules']}
            if not isinstance(schedules,list) or len(schedules)!=len(expected) or {s.get('row') for s in schedules}!=expected:raise ValueError('Danh sách khối cấu hình không khớp mẫu.')
            for schedule in schedules:
                mass=schedule.get('massDays');lesson=schedule.get('lessonDays')
                if not isinstance(mass,list) or not mass or len(set(mass))!=len(mass) or any(type(d) is not int or not 0<=d<=6 for d in mass):raise ValueError('Chọn các ngày lễ hợp lệ; Chủ Nhật=0, thứ Hai=1.')
                if 0 not in mass:raise ValueError('Giữ lịch Thánh lễ Chủ Nhật; chỉ điều chỉnh ngày lễ trong tuần.')
                if len(mass)>3:raise ValueError('Mẫu có 14 cột lễ mỗi tháng, chỉ hỗ trợ tối đa 3 ngày lễ mỗi tuần. Cần mở rộng mẫu nếu muốn thêm.')
                start_year=int(self.catalog.public()['year'][:4])
                for m in range(1,13):
                    y=start_year if m>=9 else start_year+1
                    count=sum((date(y,m,d).weekday()+1)%7 in mass for d in range(1,monthrange(y,m)[1]+1))
                    if count>14:raise ValueError(f'Lịch {y}-{m:02} có {count} buổi, vượt 14 cột lễ của mẫu. Chọn lịch khác hoặc mở rộng mẫu trước.')
                if lesson!=[0]:raise ValueError('Giáo lý chỉ vào Chủ Nhật.')
                for d in range(7):cells[f'{chr(72+d)}{schedule["row"]}']=1 if d in mass else None
                for d in range(7):cells[f'{chr(81+d)}{schedule["row"]}']=1 if d in lesson else None
            identifiers=list(self.catalog.classes) if values.get('scope')=='all' else [class_id]
            self.parity_guard(identifiers)
            payload={'optionCells':cells,'thresholdCells':thresholds}
            description=[f'Áp dụng {len(identifiers)} lớp; kiểm tra lịch, đơn và chuyên cần trước khi ghi.',
                         f'Ngưỡng Giáo Lý {thresholds["D3"]}; Thánh Lễ {thresholds["D4"]}.']
        elif operation=='note':
            content=str(values.get('content','')).strip()
            if len(content)>2000 or content.startswith('='):raise ValueError('Ghi chú tối đa 2.000 ký tự, không nhập công thức.')
            read=self.parity_read(class_id);ref=values.get('cell') or read['nextNoteCell']
            if not re.fullmatch(r'[A-G](?:[2-9]|[1-9]\d{1,2})',str(ref)):raise ValueError('Ô ghi chú phải trong vùng A2:G999.')
            old=next((n for n in read['notes'] if n['cell']==ref),None)
            if old and not old['editable']:raise ValueError('Ô ghi chú có công thức/số. Chưa hỗ trợ sửa từ web.')
            if not old and ref!=read['nextNoteCell']:raise ValueError('Chỉ thêm tại dòng trống được đề xuất.')
            payload={'cell':ref,'before':old['content'] if old else '', 'content':content}
            description=[f'Ghi vào Ghi Chú!{ref}; giữ các ghi chú còn lại.']
        elif operation=='scratch':
            mode=values.get('mode')
            if mode not in ('clear','restore'):raise ValueError('Chọn xóa hoặc khôi phục vùng nhập.')
            payload={'mode':mode}
            description=['Chỉ tác động K:L, R:S, Y:Z, hàng 4–503 của DiemDanhAuto.',
                         'Không xóa điểm, chuyên cần đã ghi hoặc nhật ký web.',
                         'Khôi phục lấy vùng mẫu CL:CM giống macro gốc.' if mode=='restore' else 'Xóa các dữ liệu đang chờ trong ba vùng nhập.']
        else:raise ValueError('Thao tác không hợp lệ.')
        revisions={cid:digest(self.catalog.classes[cid]['path']) for cid in identifiers}
        plan={'operation':operation,'classId':class_id,'classIds':identifiers,'payload':payload,
              'actor':actor,'reason':reason,'description':description,'issues':issues,'revisions':revisions}
        # Requests, scans and status changes between preview and save invalidate confirmation.
        plan['token']=token({**plan,'database':digest(self.database)})
        return plan

    def parity_preview(self,values):
        plan=self.parity_plan(values)
        if plan['operation']=='configuration' and not plan['issues']:
            for cid in plan['classIds']:
                try:_,calendar=self.parity_engine().parity_package(cid,plan['operation'],plan['payload'])
                except ValueError as error:
                    plan['issues'].append(self.catalog.classes[cid]['name']+': '+str(error));continue
                old=self.catalog.classes[cid]['calendar']
                removed={a:set(old[a])-set(calendar[a]) for a in old}
                with self.db() as db:
                    for lock in db.execute('SELECT * FROM locks WHERE class_id=?',(cid,)):
                        if lock['day'] in removed[lock['activity']]:plan['issues'].append('Lịch mới sẽ mất buổi đã chốt: '+lock['day'])
                    codes={s['id'] for s in self.catalog.students.values() if s['classId']==cid}
                    for r in db.execute("SELECT * FROM requests WHERE status IN ('pending','approved','completed')"):
                        if r['code'] in codes and r['day'] in removed[r['activity']]:plan['issues'].append('Lịch mới sẽ mất ngày của đơn: '+r['day'])
                    for r in db.execute("SELECT * FROM supplemental_mass WHERE status='active'"):
                        if r['code'] in codes and r['actual_day'] in calendar['mass']:plan['issues'].append('Lịch mới trùng ngày lễ đã cộng bổ sung: '+r['actual_day'])
                plan.setdefault('calendarChanges',[]).append({'className':self.catalog.classes[cid]['name'],
                    'removed':{a:sorted(v) for a,v in removed.items()},
                    'added':{a:sorted(set(calendar[a])-set(old[a])) for a in old}})
        return plan

    def parity_commit(self,values):
        plan=self.parity_preview(values)
        if plan['issues']:raise ValueError('; '.join(plan['issues']))
        if values.get('previewToken')!=plan['token']:raise ValueError('Dữ liệu hoặc nội dung đã thay đổi. Xem trước lại trước khi lưu.')
        packages={}
        for cid in plan['classIds']:
            packages[cid]=self.parity_engine().parity_package(cid,plan['operation'],plan['payload'])[0]
        identifier=secrets.token_hex(12);stamp=self.stamp();operation=plan['operation'];payload=plan['payload']
        def change(db):
            if operation=='student-add':db.execute('INSERT INTO issued_codes VALUES (?,?)',(payload['code'],stamp))
            if operation=='supplemental':
                if payload['action']=='add':
                    db.execute('INSERT INTO supplemental_mass VALUES (?,?,?,?,?,?,?,?)',(identifier,payload['code'],payload['actualDate'],plan['actor'],plan['reason'],'active',stamp,stamp))
                    db.execute('INSERT INTO scans VALUES (?,?,?,?,?,?,?,?)',(identifier,payload['code'],'mass',payload['actualDate'],payload['actualDate'],'supplemental:'+identifier,'supplemental',stamp))
                else:
                    db.execute("UPDATE supplemental_mass SET status='cancelled',updated_at=? WHERE id=?",(stamp,payload['id']))
                    db.execute('DELETE FROM scans WHERE request_id=?',('supplemental:'+payload['id'],))
            self.management_event(db,operation,payload.get('code',plan['classId']),payload.get('action','updated'),plan['actor'],plan['reason']+' · '+'; '.join(plan['description']))
            if operation=='supplemental':
                self.management_event(db,'supplemental',payload.get('id') or identifier,payload['action'],plan['actor'],plan['reason'])
            # SQLite grade snapshots are derived, never discard audit or unresolved jobs.
            if self.grade_mode=='portable' and not self.grade_service().repository.mysql:
                for cid in plan['classIds']:db.execute('DELETE FROM xh_grade_books WHERE class_id=?',(self.grade_service().key(cid),))
        self.parity_transaction(plan,packages,change)
        return {'saved':True,'id':identifier,'state':self.bootstrap(),'details':self.parity_read(plan['classId'])}

    def parity_transaction(self,plan,packages,change):
        directory=self.state_dir/'parity-transactions'/secrets.token_hex(12);directory.mkdir(parents=True)
        files=[]
        for cid,package in packages.items():
            info=self.catalog.classes[cid];source=info['path']
            owners=[source.with_name('~$'+source.name),source.with_name('~$'+source.name[2:])]
            if any(p.exists() for p in owners):raise ValueError('File đang mở trong Excel. Đóng file rồi xem trước lại.')
            if digest(source)!=plan['revisions'][cid]:raise ValueError('File đã sửa bên ngoài; chưa ghi đè.')
            original=directory/(cid+'.original.xlsm');prepared=directory/(cid+'.prepared.xlsm')
            shutil.copy2(source,original);prepared.write_bytes(package)
            if macro_signature(prepared)[0]!=self.catalog.trusted_signature:raise ValueError('VBA trong bản mới không khớp.')
            files.append({'path':str(source.resolve()),'original':str(original.resolve()),'prepared':str(prepared.resolve()),
                          'oldSha':plan['revisions'][cid],'newSha':digest(prepared)})
        # The SQLite transaction is prepared as a file too, so a crash can roll back both.
        original=directory/'original.sqlite3';prepared=directory/'prepared.sqlite3'
        before=digest(self.database);shutil.copy2(self.database,original);shutil.copy2(self.database,prepared)
        db=sqlite3.connect(prepared)
        try:
            with db:change(db)
        finally:db.close()
        files.append({'path':str(self.database.resolve()),'original':str(original.resolve()),'prepared':str(prepared.resolve()),'oldSha':before,'newSha':digest(prepared)})
        for f in files:
            if digest(f['path'])!=f['oldSha'] or digest(f['original'])!=f['oldSha']:raise ValueError('Dữ liệu đã thay đổi khi chuẩn bị giao dịch.')
        journal={'status':'prepared','files':files};journal_path=directory/'journal.json'
        journal_path.write_text(json.dumps(journal),encoding='utf-8')
        try:
            for f in files:
                if digest(f['path'])!=f['oldSha']:raise ValueError('Dữ liệu vừa đổi bên ngoài. Chưa ghi đè.')
                # Copy next to target so os.replace stays atomic on the target volume.
                stage=Path(f['path']).with_name('.parity-'+directory.name+Path(f['path']).suffix)
                shutil.copy2(f['prepared'],stage);os.replace(stage,f['path'])
            journal['status']='committed';journal_path.write_text(json.dumps(journal),encoding='utf-8')
        except Exception:
            for f in files:
                if digest(f['path']) not in (f['oldSha'],f['newSha']):raise ValueError('Giao dịch lỗi và file đã sửa ngoài. Giữ bản sao để kiểm tra, chưa ghi đè.')
            for f in files:
                stage=Path(f['path']).with_name('.rollback-'+directory.name+Path(f['path']).suffix)
                shutil.copy2(f['original'],stage);os.replace(stage,f['path'])
            journal['status']='rolled-back';journal_path.write_text(json.dumps(journal),encoding='utf-8')
            self.catalog.stamps={};self.catalog.refresh();raise
        self.catalog.stamps={};self.catalog.refresh()

    def bulk_preview(self,values):
        actor,reason=self.parity_manager(values)
        self.catalog.refresh()
        class_id=values.get('classId');self.parity_guard([class_id])
        info=self.catalog.classes[class_id];kind=values.get('kind');text=values.get('text')
        if values.get('expectedRevision') and values['expectedRevision']!=digest(info['path']):raise ValueError('Sổ lớp đã thay đổi. Làm mới trước khi nhập hàng loạt.')
        if not isinstance(text,str) or not text.strip() or len(text)>30000:raise ValueError('Dán danh sách từ 1 đến 40 dòng, tối đa 30.000 ký tự.')
        lines=[line.split('\t') for line in text.strip('\r\n').splitlines() if line.strip()]
        if not 1<=len(lines)<=40:raise ValueError('Mỗi lượt nhập từ 1 đến 40 dòng của một lớp.')
        entries=[];issues=[];seen=set()
        for n,line in enumerate(lines,1):
            code=line[0].strip().upper();student=self.catalog.students.get(code)
            if code in seen or not student or student['classId']!=class_id:issues.append(f'Dòng {n}: mã trùng hoặc không thuộc lớp.');continue
            seen.add(code)
            if kind=='grades':
                from validation import validate_grades
                if len(line)!=25:issues.append(f'Dòng {n}: cần 25 cột (mã + 20 khảo bài + 4 điểm kiểm tra/thi).');continue
                item={'code':code,'khaoBai':line[1:21],**dict(zip(('kt15','kt1Tiet','thiKinh','thiHk'),line[21:]))}
                try:entries.extend(validate_grades([item],class_id,self.catalog.students))
                except ValueError as e:issues.append(f'Dòng {n}: {e}')
            elif kind=='attendance':
                if len(line)!=1:issues.append(f'Dòng {n}: chỉ nhập mã; ngày và hoạt động chọn chung phía trên.');continue
                actual=day(values.get('date'));activity=values.get('activity')
                if activity not in ('mass','lesson','early'):raise ValueError('Hoạt động không hợp lệ.')
                if actual>datetime.now(self.local_tz).date().isoformat():issues.append('Không điểm danh có mặt cho ngày chưa diễn ra.');continue
                if not self.student_active(code,actual) or student['enrolledFrom'] and actual<student['enrolledFrom']:
                    issues.append(f'{code}: chưa nhập học hoặc đang nghỉ.');continue
                makeups=[r for r in self.requests() if r['code']==code and r['kind']=='makeup' and r['status']=='approved' and r['effective_target_day']==actual]
                target=makeups[0]['day'] if activity=='mass' and makeups else actual
                if not makeups and actual not in info['calendar'][activity]:issues.append(f'{code}: ngày không có trong lịch/đơn đổi lễ.');continue
                status=self.status_at(code,target,activity)
                if status in ('P','O'):issues.append(f'{code}: đã có {status}; chỉnh riêng từng em có lý do trước khi nhập hàng loạt.');continue
                if activity=='early':
                    try:self.validate_early_attendance({'date':actual,'activity':activity},{**values,'manual':True})
                    except ValueError as e:issues.append(str(e));continue
                with self.db() as db:
                    used=db.execute("SELECT status FROM scans WHERE code=? AND activity=? AND actual_day=?",(code,activity,actual)).fetchone()
                if used and used['status']=='supplemental':issues.append(f'{code}: lượt đã cộng lễ bổ sung.');continue
                if not status and self.locked({'date':target,'activity':activity},class_id):issues.append(f'{code}: buổi đã chốt; chỉnh riêng từng em.');continue
                entries.append({'code':code,'name':student['name'],'targetDate':target,'duplicate':status=='X'})
            else:raise ValueError('Loại nhập hàng loạt không hợp lệ.')
        semester=values.get('semester',1)
        if kind=='grades' and (type(semester) is not int or semester not in (1,2)):raise ValueError('Chọn học kỳ 1 hoặc 2.')
        plan={'kind':kind,'classId':class_id,'entries':entries,'issues':issues,'revision':digest(info['path']),
              'semester':semester,'date':values.get('date'),'activity':values.get('activity'),'actor':actor,'reason':reason,
              'attendedTime':values.get('attendedTime')}
        plan['token']=token({**plan,'database':digest(self.database)})
        return plan

    def bulk_commit(self,values):
        plan=self.bulk_preview(values)
        if plan['issues']:raise ValueError('; '.join(plan['issues']))
        if values.get('previewToken')!=plan['token']:raise ValueError('Danh sách đã thay đổi. Xem trước lại.')
        if plan['kind']=='grades':
            result=self.update_grades({'classId':plan['classId'],'semester':plan['semester'],'grades':plan['entries'],'expectedRevision':plan['revision']})
            with self.db() as db:self.management_event(db,'bulk-grades',plan['classId'],'updated',plan['actor'],plan['reason'])
            return {**result,'state':self.bootstrap()}
        session=self.create_session({'date':plan['date'],'activity':plan['activity'],'scope':self.catalog.classes[plan['classId']]['name']})
        completed=[]
        for item in plan['entries']:
            try:
                self.workflow_scan({**values,'code':item['code'],'sessionId':session['id'],'manual':True})
                completed.append(item['code'])
            except Exception as error:
                return {'saved':False,'partial':True,'completed':completed,'error':str(error), 'failedCode':item['code'],'state':self.bootstrap()}
        with self.db() as db:self.management_event(db,'bulk-attendance',plan['classId'],'updated',plan['actor'],plan['reason'])
        return {'saved':True,'completed':completed,'state':self.bootstrap()}

