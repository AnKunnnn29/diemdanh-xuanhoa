"""School-year isolation and reviewed transfers; preserve codes and history."""
import copy
import secrets
from datetime import datetime
from sqlalchemy import select, insert, update, delete
from openpyxl.utils.datetime import to_excel
from . import models as m
from .persistence import document, save_document, fingerprint
from .security import CloudError
from .workbooks import Workbook
from .workflows import Catalog, Workflow


class YearRules:
    def transfer_preview(self, values):
        self.reason(values)
        code, cid = str(values.get('code') or '').strip().upper(), values.get('classId')
        p, target = self.catalog.students.get(code), self.catalog.classes.get(cid)
        if not p or not target or p['classId'] == cid:
            raise CloudError('Chọn thiếu nhi và lớp đích khác lớp hiện tại.')
        self.parity_guard([cid,p['classId']])
        issues = []
        if target['count'] >= 40:
            issues.append('Lớp đích đã đủ 40 em.')
        if any(r['code'] == code and r['status'] in ('pending','approved') for r in self.requests()):
            issues.append('Xử lý hoặc dừng các đơn đang chờ/đã duyệt trước khi chuyển lớp.')
        for key in self.catalog.marks:
            day, activity, person = key.split('|')
            if person == code and day not in target['calendar'][activity]:
                issues.append(f'{day}: lớp đích không có hoạt động {activity} để giữ lịch sử.')
        plan = {'code':code, 'classId':cid, 'issues':issues,
            'description':f'{p["name"]}: {p["className"]} → {target["name"]}. Giữ mã, hồ sơ, điểm và chuyên cần.'}
        plan['token'] = fingerprint({'plan':plan,'source':self.book_revision(p['classId']),
            'target':self.book_revision(cid),'state':self.state_revision(),'reason':values['reason']})
        return plan

    def transfer(self, values):
        plan = self.transfer_preview(values)
        if plan['issues'] or values.get('previewToken') != plan['token']:
            raise CloudError('Còn xung đột hoặc dữ liệu đã thay đổi. Xem trước lại.', 409, 'revision_conflict')
        self.snapshot('Trước chuyển lớp '+plan['code'])
        code, cid = plan['code'], plan['classId']
        source = self.book(self.catalog.students[code]['classId'])
        target = self.book(cid)
        grades = copy.deepcopy(target.overrides)
        for semester in (1,2):
            row = next(r for r in source.response(semester)['students'] if r['code'] == code)
            grades.setdefault(str(semester), {})[code] = {key:row[key] for key in ('khaoBai','kt15','kt1Tiet','thiKinh','thiHk')}
        save_document(self.connection, self.year_id, 'grade-inputs', cid, grades)
        options = copy.deepcopy(target.options)
        options.setdefault('supplements', {})[code] = next(r for r in source.response(3)['students'] if r['code'] == code)['supplementalMass']
        save_document(self.connection, self.year_id, 'workbook-options', cid, options)
        self.connection.execute(update(m.students).where(m.students.c.year_id == self.year_id, m.students.c.code == code)
            .values(class_id=cid, revision=m.students.c.revision+1))
        self.management_event(self.memory, 'student', code, 'transferred', self.user['displayName'], plan['description']+' '+values['reason'])
        self.flush()
        return {'saved':True}

    def prepare_year(self, values):
        plan = self.year_preview(values)
        if plan['issues'] or values.get('previewToken') != plan['token']:
            raise CloudError('Duyệt từng em, xử lý sức chứa rồi xem trước lại trước khi tạo năm mới.', 409, 'revision_conflict')
        label = f'{plan["startYear"]}-{plan["startYear"]+1}'
        if self.connection.execute(select(m.years.c.id).where(m.years.c.label == label)).first():
            raise CloudError('Năm học này đã được chuẩn bị. Kiểm tra danh sách năm học trước khi tạo lại.')
        self.snapshot('Trước tạo năm học '+label)
        identifier = secrets.token_hex(16)
        self.connection.execute(insert(m.years).values(id=identifier, label=label, active=False, created_at=self.stamp()))
        for cid, group in self.catalog.classes.items():
            payload = {k:v for k,v in group.items() if k not in ('id','name','revision','archived')}
            payload.update(fresh=True, count=plan['counts'].get(cid,0))
            self.connection.execute(insert(m.classes).values(id=cid, year_id=identifier, name=group['name'], archived=group['archived'], payload=payload, revision=1))
            save_document(self.connection, identifier, 'source-workbook', cid,
                document(self.connection, self.year_id, 'source-workbook', cid))
            options = copy.deepcopy(self.book(cid).options)
            options.pop('supplements', None)
            for row, month in enumerate([9,10,11,12,1,2,3,4,5,6,7,8],4):
                value = datetime(plan['startYear'] if month >= 9 else plan['startYear']+1, month, 1)
                options.setdefault('cells', {}).setdefault('OPTION', {})[f'A{row}'] = to_excel(value,self.book(cid).engine.epoch)
            save_document(self.connection, identifier, 'workbook-options', cid, options)
        for placement in plan['placements']:
            if placement['classId'] == 'graduate':
                continue
            old = self.catalog.students[placement['code']]
            payload = {k:v for k,v in old.items() if k not in ('id','classId','className','revision','rosterRow')}
            payload['enrolledFrom'] = None
            self.connection.execute(insert(m.students).values(year_id=identifier, code=old['id'], class_id=placement['classId'], payload=payload, revision=1))
        assignments = [{**dict(r), 'year_id':identifier} for r in self.connection.execute(select(m.assignments)
            .where(m.assignments.c.year_id == self.year_id)).mappings()]
        if assignments:
            self.connection.execute(insert(m.assignments),assignments)
        catalog = Catalog(self.connection, identifier)
        for cid, group in catalog.classes.items():
            compiled = Workbook(self.connection, identifier, cid, catalog)
            payload = {k:v for k,v in group.items() if k not in ('id','name','revision','archived')}
            payload['calendar'] = compiled.calendar
            semesters = [compiled.engine.cell('OPTION',f'B{r}') for r in range(4,16)]
            month_order = [9,10,11,12,1,2,3,4,5,6,7,8]
            payload['semesterMonths'] = {key:sorted({day[:7] for day in compiled.calendar['mass']
                if semesters[month_order.index(int(day[5:7]))] == semester}) for key,semester in (('hk1','HKI'),('hk2','HKII'))}
            self.connection.execute(update(m.classes).where(m.classes.c.year_id == identifier,m.classes.c.id == cid).values(payload=payload))
        self.memory.execute('INSERT INTO year_plans VALUES (?,?,?,?,?,?,?,?)',
            (identifier,plan['startYear'],__import__('json').dumps(plan['placements']),self.user['displayName'],values['reason'],'prepared',self.stamp(),identifier))
        self.flush()
        return {'saved':True,'id':identifier}

    def years(self):
        return {'active':self.year_id, 'years':[{'id':r['id'],'year':r['label'], 'createdAt':r['created_at'], 'prepared':not r['active']}
            for r in self.connection.execute(select(m.years).order_by(m.years.c.label)).mappings()]}

    def year_activate_preview(self, values):
        self.reason(values)
        target = self.connection.execute(select(m.years).where(m.years.c.id == values.get('id'))).mappings().first()
        if not target or target['active']:
            raise CloudError('Chọn một năm học khác năm đang làm việc.')
        catalog = Catalog(self.connection,target['id'])
        plan = {'id':target['id'],'year':target['label'],'count':len(catalog.students),'classes':len(catalog.classes),
            'warning':'Tất cả phiên đăng nhập sẽ kết thúc. Điểm, chuyên cần, đơn và hồ sơ được giữ riêng theo từng năm.'}
        plan['token'] = fingerprint({'plan':plan, 'current':self.backup_revision(), 'reason':values['reason']})
        return plan

    def activate_year(self, values):
        plan = self.year_activate_preview(values)
        if values.get('previewToken') != plan['token'] or values.get('confirmation') != plan['year']:
            raise CloudError('Xem trước lại và nhập đúng năm học đích để xác nhận.',409,'revision_conflict')
        self.snapshot('Trước chuyển năm làm việc')
        self.connection.execute(update(m.years).values(active=False))
        self.connection.execute(update(m.years).where(m.years.c.id == plan['id']).values(active=True))
        self.connection.execute(delete(m.sessions))
        return {'saved':True,'requireReauth':True}

    def year_history(self, query):
        target = self.connection.execute(select(m.years).where(m.years.c.id == query.get('id'))).mappings().first()
        if not target:
            raise CloudError('Không tìm thấy năm học.',404,'not_found')
        workflow = Workflow(self.connection,target['id'])
        try:
            state = workflow.public({**self.user, 'classIds':list(workflow.catalog.classes)})
            code = query.get('code')
            people = [p for p in workflow.catalog.students.values() if not code or p['id'] == code]
            codes = {p['id'] for p in people}
            result = {'year':target['label'], 'students':people,
                'marks':{k:v for k,v in workflow.catalog.marks.items() if k.rsplit('|',1)[-1] in codes},
                **{key:[r for r in state[key] if r['code'] in codes] for key in ('requests','careNotes','studentStatuses')}}
            if code and people:
                book = Workbook(self.connection,target['id'],people[0]['classId'],workflow.catalog)
                result['grades'] = [{**next(r for r in book.response(sem)['students'] if r['code'] == code), 'semester':sem} for sem in (1,2,3)]
            return result
        finally:
            workflow.shutdown()
