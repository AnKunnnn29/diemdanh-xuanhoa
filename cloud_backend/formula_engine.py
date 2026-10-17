"""Evaluate the formula subset in the supplied parish template, without Excel/VBA.

Unsupported functions fail explicitly. Formula errors remain typed Excel errors;
they are never converted into a made-up zero or a successful result.
"""
import math
import re
from dataclasses import dataclass
from datetime import datetime, timezone, timedelta
from functools import lru_cache, cached_property

from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.utils.datetime import from_excel, to_excel


@dataclass(frozen=True)
class ExcelError:
    code: str


@dataclass
class Range:
    values: list
    reusable: bool = False

    @cached_property
    def flattened(self):
        return [value for row in self.values for value in row]

    def flat(self):
        return self.flattened

    @cached_property
    def groups(self):
        groups={}
        for index,value in enumerate(self.flattened):
            key=(type(value),value)
            if key not in groups:groups[key]=(value,[])
            groups[key][1].append(index)
        return list(groups.values())


class UnsupportedFormula(ValueError):
    pass


def number(value):
    if isinstance(value, ExcelError):
        return value
    if value in (None, ''):
        return 0.0
    if isinstance(value, (int,float,bool)):
        return float(value)
    try:
        return float(value)
    except (TypeError,ValueError):
        return ExcelError('#VALUE!')


def text(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, (int,float)) and float(value).is_integer():
        return str(int(value))
    return str(value)


def compare(a,b,op):
    if isinstance(a,ExcelError):return a
    if isinstance(b,ExcelError):return b
    if a is None:a='' if isinstance(b,str) else 0
    if b is None:b='' if isinstance(a,str) else 0
    # Excel compares text case-insensitively, and numbers sort below text.
    if isinstance(a,str) and isinstance(b,str):a,b=a.casefold(),b.casefold()
    if isinstance(a,str)!=isinstance(b,str):
        result=1 if isinstance(a,str) else -1
    else:result=(a>b)-(a<b)
    return {'=':result==0,'<>':result!=0,'<':result<0,'>':result>0,'<=':result<=0,'>=':result>=0}[op]


class Parser:
    precedence={'=':1,'<>':1,'<':1,'>':1,'<=':1,'>=':1,'&':2,'+':3,'-':3,'*':4,'/':4,'^':5}

    def __init__(self,formula):
        self.tokens=[t for t in Tokenizer(formula).items if t.type!='WHITE-SPACE']
        self.pos=0

    def take(self):
        t=self.tokens[self.pos];self.pos+=1;return t

    def expression(self,minimum=0):
        t=self.take()
        if t.type=='OPERATOR-PREFIX':
            left=('unary',t.value,self.expression(6))
        elif t.type=='FUNC' and t.subtype=='OPEN':
            args=[]
            if not (self.tokens[self.pos].type=='FUNC' and self.tokens[self.pos].subtype=='CLOSE'):
                while True:
                    args.append(self.expression())
                    if self.tokens[self.pos].type!='SEP':break
                    self.take()
            close=self.take()
            if close.type!='FUNC' or close.subtype!='CLOSE':raise UnsupportedFormula('Công thức thiếu dấu đóng.')
            left=('call',t.value[:-1].upper(),args)
        elif t.type=='PAREN' and t.subtype=='OPEN':
            left=self.expression();self.take()
        elif t.type=='OPERAND':
            if t.subtype=='NUMBER':left=('value',float(t.value))
            elif t.subtype=='TEXT':left=('value',t.value[1:-1].replace('""','"'))
            elif t.subtype=='LOGICAL':left=('value',t.value.upper()=='TRUE')
            elif t.subtype=='ERROR':left=('value',ExcelError(t.value))
            elif t.subtype=='RANGE':left=('ref',t.value)
            else:raise UnsupportedFormula('Loại dữ liệu công thức chưa hỗ trợ: '+t.subtype)
        else:raise UnsupportedFormula('Thành phần công thức chưa hỗ trợ: '+t.value)
        while self.pos<len(self.tokens):
            t=self.tokens[self.pos]
            if t.type=='OPERATOR-POSTFIX' and t.value=='%':
                self.take();left=('binary','/',left,('value',100.0));continue
            if t.type!='OPERATOR-INFIX' or self.precedence.get(t.value,-1)<minimum:break
            self.take();right=self.expression(self.precedence[t.value]+(0 if t.value=='^' else 1))
            left=('binary',t.value,left,right)
        return left


@lru_cache(maxsize=8192)
def parse(formula):
    parser=Parser(formula)
    result=parser.expression()
    if parser.pos!=len(parser.tokens):raise UnsupportedFormula('Công thức còn phần chưa được xử lý.')
    return result


class FormulaEngine:
    supported={'IF','OR','AND','SUM','COUNT','COUNTA','COUNTIF','SUMIF','MAX','MIN',
               'VLOOKUP','HLOOKUP','RANK','MONTH','YEAR','WEEKDAY','UPPER','NOW'}

    def __init__(self,workbook):
        self.cells={}
        self.formulas=set()
        self.bounds={}
        self.epoch=workbook.epoch
        self.cache={}
        self.ranges={}
        self.selections={}
        self.range_items=0
        self.selection_items=0
        self.active=set()
        for sheet in workbook:
            self.bounds[sheet.title]=(sheet.max_row,sheet.max_column)
            for row in sheet:
                for cell in row:
                    if cell.value is not None:
                        value=to_excel(cell.value,self.epoch) if isinstance(cell.value,datetime) else cell.value
                        self.cells[(sheet.title,cell.coordinate)]=value
                        if cell.data_type=='f':self.formulas.add((sheet.title,cell.coordinate))

    def set(self,sheet,cell,value,*,formula=False):
        key=(sheet,cell.replace('$','').upper())
        self.cells[key]=value
        self.formulas.discard(key)
        if formula:self.formulas.add(key)
        self.cache.clear()
        self.ranges.clear()
        self.selections.clear()
        self.range_items=0
        self.selection_items=0

    def cell(self,sheet,cell):
        key=(sheet,cell.replace('$','').upper())
        if key in self.cache:return self.cache[key]
        if key in self.active:raise UnsupportedFormula(f'Công thức vòng tại {sheet}!{cell}.')
        value=self.cells.get(key)
        if key in self.formulas:
            self.active.add(key)
            try:value=self.evaluate(parse(value),sheet)
            except UnsupportedFormula as error:raise UnsupportedFormula(f'{sheet}!{cell}: {error}') from error
            finally:self.active.remove(key)
        self.cache[key]=value
        return value

    def reference(self,ref,current):
        if '!' in ref:
            sheet,ref=ref.rsplit('!',1);sheet=sheet.strip("'").replace("''", "'")
        else:sheet=current
        if sheet not in self.bounds or '[' in ref:raise UnsupportedFormula('Tham chiếu ngoài mẫu: '+ref)
        ref=ref.replace('$','').upper()
        if ':' not in ref:
            if not re.fullmatch(r'[A-Z]+\d+',ref):raise UnsupportedFormula('Tham chiếu chưa hỗ trợ: '+ref)
            return self.cell(sheet,ref)
        range_key=(sheet,ref)
        if range_key in self.ranges:return self.ranges[range_key]
        a,b=ref.split(':')
        matches=[re.fullmatch(r'([A-Z]+)(\d*)',v) for v in (a,b)]
        if not all(matches):raise UnsupportedFormula('Vùng chưa hỗ trợ: '+ref)
        c1,c2=[column_index_from_string(m[1]) for m in matches]
        r1=int(matches[0][2] or 1);r2=int(matches[1][2] or self.bounds[sheet][0])
        result=Range([[self.cell(sheet,f'{get_column_letter(c)}{r}') for c in range(c1,c2+1)] for r in range(r1,r2+1)])
        # Serverless memory is bounded. Large/sliding ranges stay temporary;
        # retaining every range of the full template can exceed the RAM limit.
        size=max(0,r2-r1+1)*max(0,c2-c1+1)
        if len(self.ranges)<1000 and self.range_items+size<=50_000:
            self.ranges[range_key]=result
            self.range_items+=size
            result.reusable=True
        return result

    def evaluate(self,node,sheet):
        kind=node[0]
        if kind=='value':return node[1]
        if kind=='ref':return self.reference(node[1],sheet)
        if kind=='unary':
            value=number(self.evaluate(node[2],sheet))
            return value if isinstance(value,ExcelError) or node[1]=='+' else -value
        if kind=='binary':
            op=node[1];a=self.evaluate(node[2],sheet);b=self.evaluate(node[3],sheet)
            if isinstance(a,ExcelError):return a
            if isinstance(b,ExcelError):return b
            if op in ('=','<>','<','>','<=','>='):return compare(a,b,op)
            if op=='&':return text(a)+text(b)
            a,b=number(a),number(b)
            if isinstance(a,ExcelError):return a
            if isinstance(b,ExcelError):return b
            if op=='/' and b==0:return ExcelError('#DIV/0!')
            try:return {'+':lambda:a+b,'-':lambda:a-b,'*':lambda:a*b,'/':lambda:a/b,'^':lambda:a**b}[op]()
            except (ValueError,OverflowError):return ExcelError('#NUM!')
        name,args=node[1],node[2]
        if name not in self.supported:raise UnsupportedFormula('Hàm chưa được kiểm thử: '+name)
        if name=='IF':
            condition=self.evaluate(args[0],sheet)
            if isinstance(condition,ExcelError):return condition
            return self.evaluate(args[1 if condition else 2],sheet) if condition or len(args)>2 else False
        values=[self.evaluate(a,sheet) for a in args]
        flat=[v for a in values for v in (a.flat() if isinstance(a,Range) else [a])]
        # COUNT/COUNTA and conditional aggregation ignore errors outside matching sums.
        if name not in ('COUNT','COUNTA','COUNTIF','SUMIF','VLOOKUP','HLOOKUP'):
            error=next((v for v in flat if isinstance(v,ExcelError)),None)
            if error:return error
        numeric=[v for v in flat if isinstance(v,(int,float)) and not isinstance(v,bool)]
        if name=='SUM':return sum(numeric)
        if name=='COUNT':return len(numeric)
        if name=='COUNTA':return sum(v is not None for v in flat)
        if name=='MAX':return max(numeric,default=0)
        if name=='MIN':return min(numeric,default=0)
        if name=='OR':return any(flat)
        if name=='AND':return all(flat)
        if name=='UPPER':return text(values[0]).upper()
        if name=='NOW':return to_excel(datetime.now(timezone(timedelta(hours=7))).replace(tzinfo=None),self.epoch)
        if name in ('MONTH','YEAR','WEEKDAY'):
            serial=number(values[0])
            if isinstance(serial,ExcelError):return serial
            d=from_excel(serial,self.epoch)
            if not isinstance(d,datetime):d=datetime(1900,1,1)
            if name=='MONTH':return d.month
            if name=='YEAR':return d.year
            return (d.weekday()+1)%7+1
        if name in ('COUNTIF','SUMIF'):
            if isinstance(values[1],ExcelError):return values[1]
            source=values[0].flat() if isinstance(values[0],Range) else [values[0]]
            criterion=values[1]
            selection_key=(id(values[0]),type(criterion),criterion) if isinstance(values[0],Range) and values[0].reusable and isinstance(criterion,(str,int,float,bool,type(None))) else None
            selected=self.selections.get(selection_key) if selection_key is not None else None
            if selected is None:
                if isinstance(values[0],Range):
                    selected=sorted(i for value,indices in values[0].groups if self.matches(value,criterion) for i in indices)
                else:
                    selected=[i for i,v in enumerate(source) if self.matches(v,criterion)]
                if selection_key is not None and len(self.selections)<512 and self.selection_items+len(selected)<=50_000:
                    self.selections[selection_key]=selected
                    self.selection_items+=len(selected)
            if name=='COUNTIF':return len(selected)
            sums=(values[2].flat() if isinstance(values[2],Range) else [values[2]]) if len(values)>2 else source
            terms=[sums[i] for i in selected if i<len(sums)]
            error=next((v for v in terms if isinstance(v,ExcelError)),None)
            return error or sum(v for v in terms if isinstance(v,(int,float)) and not isinstance(v,bool))
        if name in ('VLOOKUP','HLOOKUP'):
            target,table,index=values[:3]
            if isinstance(target,ExcelError):return target
            if isinstance(index,ExcelError):return index
            if not isinstance(table,Range):return ExcelError('#VALUE!')
            rows=table.values if name=='VLOOKUP' else list(map(list,zip(*table.values)))
            index=int(number(index))-1
            approximate=len(values)<4 or bool(values[3])
            match=None
            for row in rows:
                if compare(row[0],target,'=') is True:
                    match=row;break
                if approximate and compare(row[0],target,'<=') is True:match=row
            if match is None:return ExcelError('#N/A')
            if index<0 or index>=len(match):return ExcelError('#REF!')
            return 0 if match[index] is None else match[index]
        if name=='RANK':
            value=number(values[0]);items=values[1].flat()
            numbers=[v for v in items if isinstance(v,(int,float)) and not isinstance(v,bool)]
            if not any(v==value for v in numbers):return ExcelError('#N/A')
            ascending=len(values)>2 and bool(values[2])
            return 1+sum(v<value if ascending else v>value for v in numbers)
        raise UnsupportedFormula(name)

    @staticmethod
    def matches(value,criterion):
        if isinstance(value,ExcelError):return False
        if isinstance(criterion,str):
            m=re.match(r'^(<=|>=|<>|=|<|>)(.*)$',criterion)
            if m:
                target=number(m[2]);target=m[2] if isinstance(target,ExcelError) else target
                return compare(value,target,m[1]) is True
            if '*' in criterion or '?' in criterion:
                pattern=re.escape(criterion).replace(r'\*','.*').replace(r'\?','.')
                return re.fullmatch(pattern,text(value),re.I) is not None
        return compare(value,criterion,'=') is True

    def grade_outputs(self):
        # Trigger all students: column activation and RANK depend on the entire class.
        for sheet in ('HK 1','HK 2'):
            for row in range(6,46):
                for col in range(29,34):self.cell(sheet,f'{get_column_letter(col)}{row}')
        for row in range(7,47):
            for col in range(5,15):self.cell('Tổng kết',f'{get_column_letter(col)}{row}')
        return {key:value for key,value in self.cache.items()
                if key in self.formulas}
