"""Compile a class workbook from a private template and current PostgreSQL data.

Never execute VBA. The original formula evaluator and ZIP patcher preserve the
template's formulas, formatting and VBA bytes. Editable values are overlaid by
student code, so an independently sorted worksheet cannot receive another
student's points. No durable /tmp files or Excel instance are involved.
"""
import io
from datetime import datetime
import openpyxl
from openpyxl.utils.cell import get_column_letter
from openpyxl.utils.datetime import to_excel, from_excel
from sqlalchemy import select
from . import models as m
from .formula_engine import FormulaEngine, ExcelError
from .persistence import document, save_document, fingerprint, load_blob, Download
from .security import CloudError
from .validation import validate_grades, profile_date
from .xlsm_package import patch_package, TemplateCell, remove_unused_identity_data, restrict_export_to_template, TEMPLATE_SHEET_TITLES
from .workflows import Catalog

MIME = 'application/vnd.ms-excel.sheet.macroEnabled.12'
FIELDS = {'kt15': 'Y', 'kt1Tiet': 'Z', 'thiKinh': 'AA', 'thiHk': 'AB'}


class Workbook:
    def __init__(self, db, year, class_id, catalog=None, *, options_override=None):
        self.db, self.year, self.class_id = db, year, class_id
        self.catalog = catalog or Catalog(db, year)
        self.info = self.catalog.classes.get(class_id)
        if not self.info:
            raise CloudError('Không tìm thấy lớp.', 404, 'not_found')
        source = document(db, year, 'source-workbook', class_id, {})
        self.source_digest = source.get('blobDigest') or source.get('sha256')
        self.package = load_blob(db, self.source_digest)
        book = openpyxl.load_workbook(io.BytesIO(self.package), read_only=True, data_only=False)
        try:
            self.engine = FormulaEngine(book)
        finally:
            book.close()
        self.inputs = {}
        self.overrides = document(db, year, 'grade-inputs', class_id, {})
        self.options = document(db, year, 'workbook-options', class_id, {})
        if options_override is not None:
            self.options = options_override
        self.notes = document(db, year, 'workbook-notes', class_id, {})
        self.original_grades = {}
        self.original_supplements = {}
        for sem in (1, 2):
            self.original_grades[str(sem)] = {}
            for code, row in self.rows('HK '+str(sem), 6, 45).items():
                self.original_grades[str(sem)][code] = {
                    'khaoBai': [self.engine.cell('HK '+str(sem), f'{get_column_letter(c)}{row}') for c in range(5, 25)],
                    **{key: self.engine.cell('HK '+str(sem), f'{col}{row}') for key, col in FIELDS.items()}}
        for code, row in self.rows('Tổng kết', 7, 46).items():
            self.original_supplements[code] = self.engine.cell('Tổng kết', f'O{row}')
        self.people = {code: p for code, p in self.catalog.students.items() if p['classId'] == class_id}
        marks = {key: value for key, value in self.catalog.marks.items() if key.rsplit('|', 1)[-1] in self.people}
        self.revision = fingerprint({'source': self.source_digest, 'class': self.info,
            'students': self.people, 'marks': marks, 'grades': self.overrides,
            'options': self.options, 'notes': self.notes})
        self.compile()

    def rows(self, sheet, first, last):
        result = {}
        for row in range(first, last+1):
            value = self.engine.cell(sheet, f'B{row}')
            if value not in (None, ''):
                code = str(value).strip().upper()
                if code in result:
                    raise CloudError('Mã thiếu nhi trùng trong sổ mẫu; chưa xuất file.')
                result[code] = row
        return result

    def put(self, sheet, ref, value):
        old = self.engine.cells.get((sheet, ref))
        if (sheet,ref) in self.engine.formulas:
            current = self.engine.cell(sheet, ref)
            if current == value or current in (None, '') and value in (None, ''):
                return
            raise CloudError(f'Ô {sheet}!{ref} có công thức; chưa ghi đè.')
        self.inputs[(sheet, ref)] = value
        self.engine.set(sheet, ref, value)

    def compile(self):
        # Retain each existing roster slot, then use blank slots for new codes.
        original = self.rows('Danh Sách', 7, 46)
        slots = {code: row for code, row in original.items() if code in self.people}
        for code in sorted(self.people):
            if code not in slots:
                slots[code] = next((row for row in range(7, 47) if row not in slots.values()), None)
                if slots[code] is None:
                    raise CloudError('Lớp vượt 40 em của mẫu sổ hiện tại.')
        by_row = {row: code for code, row in slots.items()}
        for row in range(7, 47):
            p = self.people.get(by_row.get(row))
            words = p['fullName'].rsplit(' ', 1) if p else ['', '']
            family, given = ('', words[0]) if len(words) == 1 else words
            values = {'B': p['id'] if p else None,
                'C': ' '.join(filter(None, (p.get('saintName'), family))) if p else None,
                'D': given if p else None,
                'G': self.info['name'].rsplit(' ', 1)[0] if p else None,
                'H': self.info['name'].rsplit(' ', 1)[-1] if p else None,
                'I': (to_excel(profile_date(p['enrolledFrom']), self.engine.epoch) if p.get('enrolledFrom') else 'All') if p else None,
                'J': p.get('gender', '') if p else None}
            for col, value in values.items():
                self.put('Danh Sách', f'{col}{row}', value)
        parts = self.info['name'].rsplit(' ', 1)
        self.put('Danh Sách', 'B2', parts[0])
        self.put('Danh Sách', 'E2', parts[1] if len(parts) == 2 else '')
        for key, ref in {'glv1':'B3', 'glv2':'E3', 'souer':'B4', 'glv3':'E4'}.items():
            self.put('Danh Sách', ref, self.info.get('teachers', {}).get(key, ''))
        for sheet, refs in self.options.get('cells', {}).items():
            for ref, value in refs.items():
                self.put(sheet, ref, value)
        for ref, value in self.notes.items():
            self.put('Ghi Chú', ref, value)
        if self.info.get('fresh'):
            for (sheet, ref), value in list(self.engine.cells.items()):
                if sheet == 'Ghi Chú' and int(''.join(c for c in ref if c.isdigit())) > 1 and (sheet, ref) not in self.inputs:
                    if not (isinstance(value, str) and value.startswith('=')):
                        self.put(sheet, ref, None)
        for row in range(5, 47):
            for col in range(5, 12):
                ref = f'{get_column_letter(col)}{row}'
                if ('Lý Lịch',ref) not in self.engine.formulas:
                    self.put('Lý Lịch', ref, None)
        for code, row in self.rows('Lý Lịch', 5, 46).items():
            if code not in self.people:
                continue
            p = self.people[code]
            for key, col in {'gender':'E', 'birthDate':'F', 'baptismDate':'G', 'parents':'H', 'address':'I', 'subParish':'J', 'phone':'K'}.items():
                value = p.get(key) or None
                if key in ('birthDate', 'baptismDate') and value:
                    value = to_excel(profile_date(value), self.engine.epoch)
                self.put('Lý Lịch', f'{col}{row}', value)
        for sem in (1, 2):
            sheet = 'HK '+str(sem)
            for row in range(6, 46):
                for col in range(5, 29):
                    self.put(sheet, f'{get_column_letter(col)}{row}', None)
            for code, row in self.rows(sheet, 6, 45).items():
                if code not in self.people:
                    continue
                base = {} if self.info.get('fresh') else self.original_grades[str(sem)].get(code, {})
                values = {**base, **self.overrides.get(str(sem), {}).get(code, {})}
                for index, value in enumerate(values.get('khaoBai', [])):
                    self.put(sheet, f'{get_column_letter(5+index)}{row}', value)
                for field, col in FIELDS.items():
                    self.put(sheet, f'{col}{row}', values.get(field))
        targets = {}
        for month in range(1, 13):
            sheet = f'T{month:02}'
            for activity, columns in {'mass':range(5, 19), 'early':range(19, 24), 'lesson':range(24, 29)}.items():
                for col in columns:
                    value = self.engine.cell(sheet, f'{get_column_letter(col)}7')
                    if isinstance(value, (int, float)) and not isinstance(value, bool) and value > 0:
                        targets[(from_excel(value, self.engine.epoch).date().isoformat(), activity)] = (sheet, col)
            for row in range(8, 48):
                for col in range(5, 29):
                    self.put(sheet, f'{get_column_letter(col)}{row}', None)
        month_rows = {f'T{month:02}': self.rows(f'T{month:02}', 8, 47) for month in range(1, 13)}
        self.calendar = {activity: sorted(day for day, a in targets if a == activity) for activity in ('mass', 'early', 'lesson')}
        for key, mark in self.catalog.marks.items():
            day, activity, code = key.split('|')
            if code not in self.people:
                continue
            if (day, activity) not in targets:
                raise CloudError('Sổ mẫu không có ngày của dữ liệu chuyên cần; chưa xuất để tránh mất lịch sử.')
            sheet, col = targets[day, activity]
            row = month_rows[sheet].get(code)
            if row is None:
                raise CloudError('Sổ mẫu chưa liên kết được mã thiếu nhi; chưa xuất file.')
            self.put(sheet, f'{get_column_letter(col)}{row}', mark['status'])
        supplements = self.options.get('supplements', {})
        for row in range(7, 47):
            code = str(self.engine.cell('Tổng kết', f'B{row}') or '').upper()
            value = supplements.get(code, None if self.info.get('fresh') else self.original_supplements.get(code))
            self.put('Tổng kết', f'O{row}', value if code in self.people else None)
        self.engine.grade_outputs()
        scratch = self.options.get('scratchMode') or ('clear' if self.info.get('fresh') else None)
        if scratch in ('clear','restore'):
            for row in range(4,504):
                for target in (11,18,25):
                    for offset,source in enumerate(('CL','CM')):
                        ref = f'{get_column_letter(target+offset)}{row}'
                        value = self.engine.cells.get(('DiemDanhAuto',f'{source}{row}')) if scratch == 'restore' else None
                        self.inputs['DiemDanhAuto',ref] = TemplateCell(value)
                        self.engine.set('DiemDanhAuto',ref,value,formula=isinstance(value,str) and value.startswith('='))

    def response(self, semester):
        if type(semester) is not int or semester not in (1, 2, 3):
            raise CloudError('Học kỳ phải là 1, 2 hoặc tổng kết năm.')
        sheet = 'HK '+str(semester) if semester < 3 else 'Tổng kết'
        first, last = (6, 45) if semester < 3 else (7, 46)
        issues, people = [], []
        def value(ref):
            result = self.engine.cell(sheet, ref)
            if isinstance(result, ExcelError):
                issues.append({'sheet': sheet, 'cell': ref, 'error': result.code})
                return result.code
            return None if result == '' else result
        for code, row in self.rows(sheet, first, last).items():
            if code not in self.people:
                continue
            p = self.people[code]
            record = {key:p.get(key) for key in ('name', 'saintName', 'fullName', 'rosterRow')}
            record['code'] = code
            if semester < 3:
                record['khaoBai'] = [value(f'{get_column_letter(c)}{row}') for c in range(5, 25)]
                record.update({key:value(f'{col}{row}') for key, col in {**FIELDS,
                    'tbHk':'AC', 'xepLoai':'AD', 'xepHang':'AE', 'massRatio':'AF', 'lessonRatio':'AG'}.items()})
            else:
                record.update({key:value(f'{col}{row}') for key, col in {'diemHk1':'E', 'diemHk2':'F', 'diemGiaoLy':'G',
                    'xepLoai':'H', 'massScore':'I', 'lessonScore':'J', 'combinedMassScore':'K',
                    'assessment':'L', 'totalScore':'M', 'yearRank':'N', 'supplementalMass':'O'}.items()})
            people.append(record)
        return {'classId': self.class_id, 'className': self.info['name'], 'semester': semester,
            'students': people, 'revision': self.revision, 'calculatedBy': 'server',
            'calculationIssues': issues, 'excelReady': True,
            'thresholds': {'catechism':self.engine.cell('Tổng kết', 'D3'), 'attendance':self.engine.cell('Tổng kết', 'D4')}}

    def save(self, semester, grades, expected):
        if semester not in (1, 2) or type(semester) is not int:
            raise CloudError('Chỉ nhập điểm học kỳ 1 hoặc 2.')
        if self.info['archived']:
            raise CloudError('Lớp đã lưu trữ; chưa thay đổi điểm.')
        if expected != self.revision:
            raise CloudError('Bảng điểm đã được sửa ở phiên khác. Tải lại trước khi lưu.', 409, 'revision_conflict')
        try:
            normalized = validate_grades(grades, self.class_id, self.catalog.students)
        except ValueError as error:
            raise CloudError(str(error)) from error
        data = {key:dict(value) for key, value in self.overrides.items()}
        records = data.setdefault(str(semester), {})
        for item in normalized:
            records[item['code']] = {**records.get(item['code'], {}), **{k:v for k,v in item.items() if k != 'code'}}
        save_document(self.db, self.year, 'grade-inputs', self.class_id, data)
        fresh = Workbook(self.db, self.year, self.class_id, self.catalog)
        return {'saved': True, 'excelReady': True, 'vbaExecuted': False, 'calculatedBy': 'server',
            'grades': fresh.response(semester), 'revision': fresh.revision, 'message': 'Đã lưu điểm trên cloud; sổ tải về được tính từ dữ liệu mới nhất.'}

    def export(self):
        # Also refresh identity caches in profile/month sheets before exporting.
        for sheet, first, last in [('Lý Lịch', 5, 46)]+[(f'T{month:02}', 8, 47) for month in range(1, 13)]:
            for row in range(first, last+1):
                for col in ('B', 'C', 'D'):
                    self.engine.cell(sheet, f'{col}{row}')
        for (sheet,ref),value in list(self.engine.cells.items()):
            if sheet in TEMPLATE_SHEET_TITLES and sheet!='Ghi Chú' and (sheet,ref) in self.engine.formulas:
                self.engine.cell(sheet,ref)
        caches = {key:value for key,value in self.engine.cache.items()
            if key in self.engine.formulas}
        package = patch_package(self.package, self.inputs, caches)
        # File xuất chỉ theo mẫu đã kiểm chứng; dữ liệu lớp khác không đi theo file GLV.
        foreign_people=[person for person in self.catalog.students.values() if person['classId']!=self.class_id]
        package = restrict_export_to_template(package, [person['id'] for person in foreign_people],
            trusted_notes=self.notes, foreign_profiles=foreign_people, own_profiles=self.people.values())
        package = remove_unused_identity_data(package, {code:p.get('saintName','') for code,p in self.people.items()},self.class_id)
        return Download(package, self.info['name']+'.xlsm', MIME)
