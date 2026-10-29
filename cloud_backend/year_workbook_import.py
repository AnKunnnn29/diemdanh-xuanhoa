"""Admin-reviewed new school years from existing XLSM templates; never execute VBA."""
import base64
import hashlib
import io
import re
import secrets
import time
import unicodedata
import zipfile
from datetime import datetime
import openpyxl
from openpyxl.utils.cell import get_column_letter
from openpyxl.utils.datetime import from_excel
from sqlalchemy import select, insert
from . import models as m
from .formula_engine import FormulaEngine, ExcelError
from .persistence import document, save_document, store_blob, load_blob, fingerprint
from .security import CloudError
from .validation import profile_date
from .xlsm_package import TEMPLATE_SHEET_TITLES, vba_hashes
from .workbooks import Workbook
from .workflows import Catalog

MAX_FILE=1_900_000
REQUIRED={'Danh Sách','Lý Lịch','HK 1','HK 2','Tổng kết','OPTION','Ghi Chú','DiemDanhAuto'}|{f'T{n:02}' for n in range(1,13)}
def clean(v):return ' '.join(unicodedata.normalize('NFC',str(v or '')).split())

def parse_workbook(content,year,known):
 try:
  with zipfile.ZipFile(io.BytesIO(content)) as archive:
   entries=archive.infolist()
   if len(entries)>2000 or len({e.filename for e in entries})!=len(entries) or sum(e.file_size for e in entries)>40_000_000 or any(e.file_size>10_000_000 or e.flag_bits&1 for e in entries):
    raise CloudError('Sổ Excel vượt giới hạn hoặc có thành phần không hợp lệ.')
   for e in entries:
    if e.filename.endswith('.xml') and b'<!DOCTYPE' in archive.read(e).upper():raise CloudError('Sổ Excel có XML không hợp lệ.')
  book=openpyxl.load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
  try:
   if not REQUIRED.issubset(book.sheetnames) or set(book.sheetnames)-TEMPLATE_SHEET_TITLES:raise CloudError('Cần sổ .xlsm đúng mẫu các lớp đang sử dụng, đủ các trang danh sách, lý lịch, điểm và 12 tháng.')
   if sum(s.max_row*s.max_column for s in book)>300_000 or any(s.max_row>2500 or s.max_column>200 for s in book):raise CloudError('Vùng dữ liệu sổ Excel quá lớn so với mẫu.')
   engine=FormulaEngine(book)
   props={p.name:p.value for p in book.custom_doc_props}
  finally:book.close()
  def value(sheet,ref):
   v=engine.cell(sheet,ref)
   if isinstance(v,ExcelError):raise CloudError('Sổ có lỗi công thức ở '+sheet+'!'+ref+'. Sửa trên Excel rồi tải lại.')
   return v
  def day(v):
   if v in (None,'','All'):return None
   if type(v) in (int,float):return from_excel(v,engine.epoch).date().isoformat()
   parsed=profile_date(str(v));return parsed.date().isoformat() if parsed else None
  name=clean(str(value('Danh Sách','B2') or '')+' '+str(value('Danh Sách','E2') or ''))
  if not 2<=len(name)<=100:raise CloudError('Chưa có tên lớp hợp lệ ở Danh Sách!B2 và E2.')
  months=[9,10,11,12,1,2,3,4,5,6,7,8];semesters={'hk1':[],'hk2':[]}
  for row,month in enumerate(months,4):
   expected=f'{year if month>=9 else year+1}-{month:02}'
   actual=day(value('OPTION',f'A{row}'))
   if not actual or actual[:7]!=expected:raise CloudError(f'Sổ {name}: lịch OPTION chưa thuộc năm {year}-{year+1}. Cập nhật lịch trong Excel trước khi tải lên.')
   sem=clean(value('OPTION',f'B{row}')).upper()
   if sem not in ('HKI','HKII'):raise CloudError('Chưa xác định học kỳ ở OPTION!B'+str(row)+'.')
   semesters['hk1' if sem=='HKI' else 'hk2'].append(expected)
  people=[];codes=set();profiles={}
  for row in range(5,45):
   code=clean(value('Lý Lịch',f'B{row}')).upper()
   if code:profiles[code]=row
  for row in range(7,47):
   code=clean(value('Danh Sách',f'B{row}')).upper();family=clean(value('Danh Sách',f'C{row}'));given=clean(value('Danh Sách',f'D{row}'))
   if not code:
    if family or given:raise CloudError(f'Sổ {name}: dòng {row} có tên nhưng thiếu mã thiếu nhi.')
    continue
   if not re.fullmatch(r'XH[0-9]{1,30}',code) or code in codes:raise CloudError(f'Sổ {name}: mã thiếu nhi sai định dạng hoặc bị trùng ở dòng {row}.')
   codes.add(code);fullname=clean(family+' '+given)
   if not 2<=len(fullname)<=200:raise CloudError(f'Sổ {name}: thiếu họ tên ở dòng {row}.')
   saint=clean(props.get('XH_Saint_'+code,''));saint='' if saint=='__XH_EMPTY__' else saint
   if not saint and code in known and clean(known[code].get('name'))==fullname:saint=known[code].get('saintName','')
   if saint and not (fullname==saint or fullname.startswith(saint+' ')):saint=''
   full=fullname[len(saint):].strip() if saint else fullname
   person={'id':code,'saintName':saint,'fullName':full,'name':fullname,'rosterRow':row,'enrolledFrom':day(value('Danh Sách',f'I{row}'))}
   profile=profiles.get(code)
   for key,col in {'gender':'E','birthDate':'F','baptismDate':'G','parents':'H','address':'I','subParish':'J','phone':'K'}.items():
    raw=value('Lý Lịch',f'{col}{profile}') if profile else None
    person[key]=(day(raw) or '') if key in ('birthDate','baptismDate') else clean(raw)
   if not person['gender']:person['gender']=clean(value('Danh Sách',f'J{row}'))
   people.append(person)
  if any(value('Danh Sách',f'{col}{r}') not in (None,'') for r in range(47,52) for col in ('B','C','D')):raise CloudError('Mỗi sổ lớp chỉ hỗ trợ 40 em trong vùng danh sách của mẫu.')
  for sheet,first,last in [('Lý Lịch',5,44),('HK 1',6,45),('HK 2',6,45),('Tổng kết',7,46)]+[(f'T{n:02}',8,47) for n in range(1,13)]:
   found=[clean(value(sheet,f'B{r}')).upper() for r in range(first,last+1) if value(sheet,f'B{r}') not in (None,'')]
   if len(found)!=len(set(found)) or set(found)!=codes:raise CloudError(f'Sổ {name}: mã thiếu nhi ở trang {sheet} chưa khớp Danh Sách. Cập nhật sổ bằng Excel rồi tải lại.')
  calendar={a:[] for a in ('mass','early','lesson')};marks=[]
  for n in range(1,13):
   sheet=f'T{n:02}'
   for activity,columns in [('mass',range(5,19)),('early',range(19,24)),('lesson',range(24,29))]:
    for col in columns:
     ref=get_column_letter(col);raw=value(sheet,ref+'7');date=day(raw) if type(raw) in (int,float) and raw>0 else None
     if not date:continue
     if not f'{year}-09-01'<=date<=f'{year+1}-08-31':raise CloudError('Lịch hoạt động không thuộc năm học đã chọn.')
     if date in calendar[activity]:raise CloudError('Ngày hoạt động bị trùng trong sổ '+name+'.')
     calendar[activity].append(date)
     for r in range(8,48):
      status=clean(value(sheet,f'{ref}{r}')).upper();code=clean(value(sheet,f'B{r}')).upper()
      if status and code in codes:
       if status not in ('X','P','O'):raise CloudError(f'Sổ {name}: ký hiệu điểm danh chưa hợp lệ ở {sheet}!{ref}{r}.')
       marks.append({'day':date,'activity':activity,'code':code,'status':status})
  teachers={key:clean(value('Danh Sách',ref)) for key,ref in {'glv1':'B3','glv2':'E3','souer':'B4','glv3':'E4'}.items()}
  return {'name':name,'students':people,'calendar':{k:sorted(v) for k,v in calendar.items()},'semesterMonths':semesters,'teachers':teachers,'marks':marks}
 except CloudError:raise
 except Exception as error:
  raise CloudError('Không đọc được sổ .xlsm theo mẫu hoặc công thức chưa được hỗ trợ. Mở, kiểm tra và lưu lại bằng Excel.') from error

class WorkbookYearImport:
 def upload_year(self,values):
  year=values.get('startYear');current=int(self.year['label'][:4])
  if type(year) is not int or not current<year<2100:raise CloudError('Chọn năm bắt đầu sau năm học hiện tại, đến 2099.')
  filename=values.get('filename','')
  if not isinstance(filename,str) or len(filename)>160 or not filename.lower().endswith('.xlsm') or '/' in filename or chr(92) in filename:raise CloudError('Chọn file sổ lớp .xlsm.')
  try:content=base64.b64decode(values.get('fileBase64',''),validate=True)
  except Exception as error:raise CloudError('Nội dung file tải lên chưa hợp lệ.') from error
  if not 0<len(content)<=MAX_FILE:raise CloudError('Mỗi sổ .xlsm tối đa 1,9 MB; tải từng sổ một lượt.')
  history=self.connection.execute(select(m.students).join(m.years,m.students.c.year_id==m.years.c.id).order_by(m.years.c.label.desc())).mappings().all()
  known={}
  for entry in history:known.setdefault(entry['code'],{'id':entry['code'],**entry['payload']})
  parsed=parse_workbook(content,year,known)
  for person in parsed['students']:
   if person['id'] in known and clean(person['name']).casefold()!=clean(known[person['id']].get('name')).casefold():raise CloudError('Mã '+person['id']+' có tên không khớp hồ sơ đã có. Kiểm tra mã trong Excel hoặc cập nhật đúng hồ sơ trước khi nhập năm mới.')
  existing=next((c for c in self.catalog.classes.values() if clean(c['name']).casefold()==parsed['name'].casefold()),None)
  template=(existing or next(iter(self.catalog.classes.values()),None))
  if not template:raise CloudError('Chưa có sổ mẫu để đối chiếu.')
  source=document(self.connection,self.year_id,'source-workbook',template['id'],{})
  trusted=load_blob(self.connection,source.get('blobDigest') or source.get('sha256'))
  if vba_hashes(content)!=vba_hashes(trusted):raise CloudError('Macro trong sổ khác mẫu đang dùng. Chưa nhận file này; cần kiểm tra mẫu trước.')
  cid=existing['id'] if existing else hashlib.sha256(parsed['name'].encode()).hexdigest()[:16]
  identifier=secrets.token_hex(16);digest=store_blob(self.connection,content)
  save_document(self.connection,self.year_id,'year-upload',identifier,{'userId':self.user['id'],'startYear':year,'uploadedAt':time.time(),'sha256':digest,'classId':cid,'filename':filename,'parsed':parsed})
  return {'id':identifier,'classId':cid,'className':parsed['name'],'filename':filename,'count':len(parsed['students']),'newStudents':sum(p['id'] not in self.catalog.students for p in parsed['students']),'marks':len(parsed['marks'])}

 def uploaded_year_preview(self,values):
  year=values.get('startYear');current=int(self.year['label'][:4]);ids=values.get('uploads')
  if type(year) is not int or not current<year<2100 or not isinstance(ids,list) or not 1<=len(ids)<=60 or any(not isinstance(i,str) for i in ids) or len(set(ids))!=len(ids):raise CloudError('Chọn năm học và tải lên từ 1 đến 60 sổ lớp.')
  label=f'{year}-{year+1}'
  if self.connection.execute(select(m.years.c.id).where(m.years.c.label==label)).first():raise CloudError('Năm học này đã có dữ liệu. Chưa ghi đè năm đã chuẩn bị.')
  entries=[];classes=set();codes=set();names=set()
  for identifier in ids:
   item=document(self.connection,self.year_id,'year-upload',identifier)
   if not item or item['userId']!=self.user['id'] or item['startYear']!=year or time.time()-item['uploadedAt']>86400:raise CloudError('File xem trước đã hết hạn hoặc không thuộc lần chuẩn bị này. Tải lại các sổ lớp.')
   parsed=item['parsed'];name=clean(parsed['name']).casefold()
   if item['classId'] in classes or name in names:raise CloudError('Có hai file cùng một lớp. Bỏ file trùng rồi xem trước lại.')
   classes.add(item['classId']);names.add(name)
   for p in parsed['students']:
    if p['id'] in codes:raise CloudError('Có mã thiếu nhi xuất hiện ở nhiều sổ lớp. Sửa Excel rồi tải lại.')
    codes.add(p['id'])
   entries.append(item)
  omitted=[{'code':p['id'],'name':p['name'],'className':p['className']} for p in self.catalog.students.values() if p['id'] not in codes]
  summary=[{'className':i['parsed']['name'],'count':len(i['parsed']['students']),'filename':i['filename'],'marks':len(i['parsed']['marks'])} for i in entries]
  token=fingerprint({'year':year,'uploads':sorted(ids),'files':sorted(i['sha256'] for i in entries),'current':self.state_revision(),'catalog':{'students':self.catalog.students,'classes':self.catalog.classes,'marks':self.catalog.marks}})
  return {'year':label,'classes':summary,'count':len(codes),'omitted':omitted,'token':token,'assignmentsCopied':False}

 def import_uploaded_year(self,values):
  self.reason(values);plan=self.uploaded_year_preview(values)
  if values.get('previewToken')!=plan['token']:raise CloudError('Dữ liệu đã thay đổi. Xem trước lại trước khi tạo năm học.',409,'revision_conflict')
  if plan['omitted'] and values.get('acceptOmitted') is not True:raise CloudError('Kiểm tra và xác nhận các em không có trong sổ năm mới.')
  entries=[document(self.connection,self.year_id,'year-upload',i) for i in values['uploads']]
  self.snapshot('Trước nhập sổ năm học '+plan['year'])
  identifier=secrets.token_hex(16)
  self.connection.execute(insert(m.years).values(id=identifier,label=plan['year'],active=False,created_at=self.stamp()))
  for item in entries:
   parsed=item['parsed'];cid=item['classId']
   payload={k:parsed[k] for k in ('calendar','semesterMonths','teachers')};payload.update(count=len(parsed['students']),filename=item['filename'])
   self.connection.execute(insert(m.classes).values(year_id=identifier,id=cid,name=parsed['name'],archived=False,payload=payload,revision=1))
   save_document(self.connection,identifier,'source-workbook',cid,{'sha256':item['sha256']})
   for person in parsed['students']:
    payload={**person,'className':parsed['name']};code=payload.pop('id')
    self.connection.execute(insert(m.students).values(year_id=identifier,code=code,class_id=cid,payload=payload,revision=1))
   for mark in parsed['marks']:
    self.connection.execute(insert(m.marks).values(year_id=identifier,**mark,source='Excel',stamp=None))
  catalog=Catalog(self.connection,identifier)
  for cid in catalog.classes:
   # Validate that imported sources can run and export in the existing application.
   book=Workbook(self.connection,identifier,cid,catalog)
   for sem in (1,2,3):
    if book.response(sem)['calculationIssues']:raise CloudError('Sổ năm mới còn lỗi công thức. Sửa file và tải lại; chưa tạo năm học.')
  self.management_event(self.memory,'year',identifier,'imported-workbooks',self.user['displayName'],values['reason'])
  save_document(self.connection,identifier,'year-import','v1',{'sourceYear':self.year_id,'uploads':values['uploads'],'createdBy':self.user['id'],'count':plan['count']})
  return {'saved':True,'id':identifier,'year':plan['year'],'classes':len(entries),'count':plan['count'],'activeChanged':False}
